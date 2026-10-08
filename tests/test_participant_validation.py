"""Rechazo de grupos ambiguos antes de publicar el programa."""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from meeting_generator.generator import generate_document
from meeting_generator.models import MeetingWeek
from meeting_generator.parser import ProgramParser, parse_document
from meeting_generator.utils import split_names
from meeting_generator.validation import GenerationValidationError
from tests.source_fixtures import build_primary_source
from tests.template_fixture import build_template


class SplitNamesValidationTests(unittest.TestCase):
    def test_valid_names_preserve_normalization_and_suffixes(self) -> None:
        cases = (
            ("  ANA   EJEMPLO  ", ("Ana Ejemplo", "")),
            ("Juan McDONALD", ("Juan McDONALD", "")),
            (
                "\tÁNGELA   EJEMPLO (h)\n//  JOSÉ EJEMPLO (p) ",
                ("Ángela Ejemplo (h)", "José Ejemplo (p)"),
            ),
            ("Ana Ejemplo//Luis Ejemplo", ("Ana Ejemplo", "Luis Ejemplo")),
            ("Ana O'Neil // Luis Pérez-Gómez", ("Ana O'Neil", "Luis Pérez-Gómez")),
        )
        for text, expected in cases:
            with self.subTest(text=text):
                self.assertEqual(split_names(text), expected)

    def test_missing_field_remains_an_empty_pair(self) -> None:
        for text in ("", " \t\n\r ", "\u00a0"):
            with self.subTest(text=text):
                self.assertEqual(split_names(text), ("", ""))

    def test_rejects_extra_separators_and_empty_groups(self) -> None:
        for text in (
            "Persona Uno // Persona Dos // Persona Tres",
            "Persona Uno // // Persona Dos",
            "// Persona Dos",
            "Persona Uno //",
            "Persona Uno // \t\n ",
            " \t\n // Persona Dos",
            "//",
            "////",
            "Persona Uno///Persona Dos",
            "Persona Uno////Persona Dos",
            "Persona Uno/Persona Dos",
            "Persona Uno / / Persona Dos",
            "Persona Uno // Persona Dos/Persona Tres",
            "(h)",
            "Persona Uno // (p)",
        ):
            with self.subTest(text=text), self.assertRaises(ValueError):
                split_names(text)


class SourceParticipantValidationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        fixtures = tempfile.TemporaryDirectory()
        cls.addClassCleanup(fixtures.cleanup)
        directory = Path(fixtures.name)
        cls.source_path = build_primary_source(directory / "source.docx")
        cls.template_path = build_template(directory / "template.docx")

    def setUp(self) -> None:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.directory = Path(directory.name)
        self.invalid_source = self.directory / "invalid-source.docx"

    def test_numbered_assignments_reject_groups_with_week_and_row(self) -> None:
        cases = (
            ("Práctica sintética 1", "Persona Uno // Persona Dos // Persona Tres"),
            ("Tema sintético de tesoros 1", "// Persona Dos"),
            ("Tema sintético de vida cristiana 1", "Persona Uno/Persona Dos"),
            ("Estudio bíblico de la congregación", "Persona Uno // "),
            ("PALABRAS DE CONCLUSIÓN", "Persona Uno // // Persona Dos"),
        )
        for title, text in cases:
            with self.subTest(title=title):
                document = Document(self.source_path)
                row_number, row = next(
                    (number, row)
                    for number, row in enumerate(document.tables[2].rows, start=1)
                    if row.cells[1].text.startswith(title)
                )
                row.cells[2].text = text
                document.save(self.invalid_source)

                with self.assertRaises(GenerationValidationError) as raised:
                    parse_document(self.invalid_source)

                self.assertTrue(
                    any(
                        f"semana 3, fila {row_number}" in issue
                        for issue in raised.exception.issues
                    ),
                    raised.exception.issues,
                )

    def test_combined_row_validates_each_participant_paragraph(self) -> None:
        for paragraph_index in (0, 1):
            for text in (
                "Persona Uno // Persona Dos // Persona Tres",
                "Persona Uno // ",
            ):
                with self.subTest(paragraph=paragraph_index, text=text):
                    document = Document(self.source_path)
                    row_number, row = next(
                        (number, row)
                        for number, row in enumerate(document.tables[7].rows, start=1)
                        if len(row.cells[0].paragraphs) == 2
                        and row.cells[1].text.startswith("Tema sintético")
                    )
                    row.cells[2].paragraphs[paragraph_index].text = text
                    document.save(self.invalid_source)

                    with self.assertRaises(GenerationValidationError) as raised:
                        parse_document(self.invalid_source)

                    self.assertTrue(
                        any(
                            f"semana 8, fila {row_number}" in issue
                            for issue in raised.exception.issues
                        ),
                        raised.exception.issues,
                    )

    def test_closing_prayer_rejects_ambiguous_and_multiple_names(self) -> None:
        for text in (
            "Persona Uno // Persona Dos // Persona Tres",
            "Persona Uno // Persona Dos",
            "Persona Uno // ",
            "//",
            "Persona Uno (h // Persona Dos)",
            "Persona Uno (h/)",
        ):
            with self.subTest(text=text):
                document = Document(self.source_path)
                row_number, row = next(
                    (number, row)
                    for number, row in enumerate(document.tables[1].rows, start=1)
                    if row.cells[1].text == "ORACIÓN FINAL"
                )
                row.cells[2].text = text
                document.save(self.invalid_source)

                with self.assertRaises(GenerationValidationError) as raised:
                    parse_document(self.invalid_source)

                self.assertTrue(
                    any(
                        f"semana 2, fila {row_number}" in issue
                        for issue in raised.exception.issues
                    ),
                    raised.exception.issues,
                )

    def test_president_rejects_ambiguous_and_multiple_names_with_week(self) -> None:
        for text in (
            "Persona Uno // Persona Dos // Persona Tres",
            "Persona Uno // Persona Dos",
            "Persona Uno // ",
            "Persona Uno/Persona Dos",
            "Persona Uno (h // Persona Dos)",
            "Persona Uno (h/)",
        ):
            with self.subTest(text=text):
                document = Document(self.source_path)
                presidents = [
                    paragraph
                    for paragraph in document.paragraphs
                    if paragraph.text.startswith("PRESIDENTE:")
                ]
                presidents[2].text = f"PRESIDENTE: {text}. CANCIÓN 44"
                document.save(self.invalid_source)

                with self.assertRaises(GenerationValidationError) as raised:
                    parse_document(self.invalid_source)

                self.assertTrue(
                    any(
                        "semana 3" in issue and "presidente" in issue.lower()
                        for issue in raised.exception.issues
                    ),
                    raised.exception.issues,
                )

    def test_closing_fallback_validates_before_removing_suffixes(self) -> None:
        parser = ProgramParser(self.source_path)
        document = Document()
        table = document.add_table(rows=2, cols=3)
        table.rows[0].cells[1].text = "CANCIÓN 61"
        table.rows[1].cells[1].text = "ORACIÓN FINAL"
        for previous_prayer in ("", "Persona Válida"):
            for text in (
                "Persona Uno // Persona Dos // Persona Tres",
                "Persona Uno (h // Persona Dos)",
                "Persona Uno (h/)",
                "Persona Uno // ",
            ):
                with self.subTest(text=text, previous_prayer=previous_prayer):
                    week = MeetingWeek(week_index=4, closing_prayer=previous_prayer)
                    table.rows[1].cells[2].text = text

                    with self.assertRaises(GenerationValidationError) as raised:
                        parser._extract_closing_from_table(week, table)

                    self.assertTrue(
                        any(
                            "semana 4, fila 2" in issue
                            for issue in raised.exception.issues
                        ),
                        raised.exception.issues,
                    )
                    self.assertEqual(week.closing_prayer, previous_prayer)

    def test_rejected_groups_do_not_call_writer_or_change_output(self) -> None:
        previous = b"programa anterior"
        output_path = self.directory / "output.docx"
        for text in (
            "Persona Uno // Persona Dos // Persona Tres",
            "Persona Uno // ",
            "//",
            "Persona Uno///Persona Dos",
        ):
            document = Document(self.source_path)
            row_number, row = next(
                (number, row)
                for number, row in enumerate(document.tables[0].rows, start=1)
                if row.cells[1].text.startswith("Práctica sintética 1")
            )
            row.cells[2].text = text
            document.save(self.invalid_source)
            for preexisting in (False, True):
                with self.subTest(text=text, preexisting=preexisting):
                    output_path.unlink(missing_ok=True)
                    if preexisting:
                        output_path.write_bytes(previous)
                    before = set(self.directory.iterdir())

                    with patch("meeting_generator.generator.fill_template") as writer:
                        result = generate_document(
                            self.invalid_source, self.template_path, output_path
                        )

                    writer.assert_not_called()
                    self.assertFalse(result.succeeded)
                    self.assertFalse(result.published)
                    self.assertEqual(result.written_weeks, 0)
                    self.assertEqual(result.omitted_weeks, 9)
                    self.assertIsNone(result.output_path)
                    self.assertTrue(
                        any(
                            f"semana 1, fila {row_number}" in error
                            for error in result.errors
                        ),
                        result.errors,
                    )
                    self.assertEqual(set(self.directory.iterdir()), before)
                    if preexisting:
                        self.assertEqual(output_path.read_bytes(), previous)
                    else:
                        self.assertFalse(output_path.exists())
