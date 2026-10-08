"""Regresiones del resultado real antes y después de publicar el DOCX."""

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from meeting_generator.generator import generate_document
from meeting_generator.parser import parse_document
from meeting_generator.results import GenerationResult
from meeting_generator.template_writer import TemplateWriter, fill_template
from tests.source_fixtures import build_primary_source
from tests.template_fixture import build_template


class PublishedResultTests(unittest.TestCase):
    def test_published_failure_preserves_counts_path_and_errors(self) -> None:
        path = Path("published.docx")

        result = GenerationResult.published_failure(
            9, path, ["  lectura fallida  ", " "]
        )

        self.assertTrue(result.published)
        self.assertFalse(result.succeeded)
        self.assertEqual(result.written_weeks, 9)
        self.assertEqual(result.omitted_weeks, 0)
        self.assertEqual(result.requested_weeks, 9)
        self.assertEqual(result.output_path, path)
        self.assertEqual(result.errors, ("lectura fallida",))

    def test_published_failure_supplies_detail_for_empty_errors(self) -> None:
        result = GenerationResult.published_failure(1, Path("published.docx"), [" "])

        self.assertTrue(result.errors)
        self.assertTrue(result.published)
        self.assertFalse(result.succeeded)

    def test_published_state_remains_true_when_output_becomes_inaccessible(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "published.docx"
            output_path.write_bytes(b"publicado")
            result = GenerationResult.success(1, output_path)
            self.assertTrue(result.succeeded)

            output_path.unlink()

            self.assertTrue(result.published)
            self.assertFalse(result.succeeded)

    def test_unpublished_failure_has_no_output(self) -> None:
        result = GenerationResult.failure(9, ["no se pudo reemplazar"])

        self.assertFalse(result.published)
        self.assertFalse(result.succeeded)
        self.assertEqual(result.written_weeks, 0)
        self.assertEqual(result.omitted_weeks, 9)

    def test_result_rejects_inconsistent_states(self) -> None:
        path = Path("published.docx")
        invalid_states = (
            (-1, 0, ("error",), None),
            (0, -1, ("error",), None),
            (1, 0, ("error",), None),
            (0, 1, (), None),
            (0, 0, ("error",), path),
            (1, 1, ("error",), path),
            (1, 1, (), path),
        )
        for values in invalid_states:
            with self.subTest(values=values), self.assertRaises(ValueError):
                GenerationResult(*values)


class AtomicPublicationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.fixtures = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.fixtures.cleanup)
        fixture_directory = Path(cls.fixtures.name)
        cls.source_path = build_primary_source(fixture_directory / "source.docx")
        cls.template_path = build_template(fixture_directory / "template.docx")
        cls.weeks = parse_document(cls.source_path)
        congregation_patcher = patch(
            "meeting_generator.template_writer.CONGREGATION_NAME",
            "Congregación de Prueba",
        )
        congregation_patcher.start()
        cls.addClassCleanup(congregation_patcher.stop)

    def setUp(self) -> None:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.output_path = Path(directory.name) / "published.docx"

    def _assert_unpublished(self, result: GenerationResult, count: int = 1) -> None:
        self.assertFalse(result.published)
        self.assertFalse(result.succeeded)
        self.assertEqual(result.written_weeks, 0)
        self.assertEqual(result.omitted_weeks, count)
        self.assertIsNone(result.output_path)
        self.assertTrue(result.errors)

    def _assert_published_error(self, result: GenerationResult, count: int = 1) -> None:
        self.assertTrue(result.published)
        self.assertFalse(result.succeeded)
        self.assertEqual(result.written_weeks, count)
        self.assertEqual(result.omitted_weeks, 0)
        self.assertEqual(result.output_path, self.output_path)
        self.assertTrue(result.errors)
        self.assertTrue(Document(self.output_path).tables)

    def test_prepublication_failures_preserve_existing_output_or_create_nothing(
        self,
    ) -> None:
        previous = b"archivo anterior"
        for operation in ("save", "digest", "replace"):
            for preexisting in (False, True):
                with self.subTest(operation=operation, preexisting=preexisting):
                    self.output_path.unlink(missing_ok=True)
                    if preexisting:
                        self.output_path.write_bytes(previous)
                    before = set(self.output_path.parent.iterdir())
                    target = {
                        "save": "docx.document.Document.save",
                        "digest": "meeting_generator.template_writer.TemplateWriter._file_digest",
                        "replace": "meeting_generator.template_writer.os.replace",
                    }[operation]

                    with patch(target, side_effect=OSError(f"fallo de {operation}")):
                        result = fill_template(
                            self.template_path, self.weeks[:1], self.output_path
                        )

                    self._assert_unpublished(result)
                    self.assertEqual(set(self.output_path.parent.iterdir()), before)
                    if preexisting:
                        self.assertEqual(self.output_path.read_bytes(), previous)
                    else:
                        self.assertFalse(self.output_path.exists())

    def test_output_read_failure_preserves_published_result(self) -> None:
        original_digest = TemplateWriter._file_digest

        def fail_output_digest(writer, path):
            if path == self.output_path:
                raise OSError("lectura posterior fallida")
            return original_digest(writer, path)

        for preexisting in (False, True):
            with self.subTest(preexisting=preexisting):
                self.output_path.unlink(missing_ok=True)
                if preexisting:
                    self.output_path.write_bytes(b"archivo anterior")
                with patch.object(TemplateWriter, "_file_digest", fail_output_digest):
                    result = fill_template(
                        self.template_path, self.weeks[:1], self.output_path
                    )

                self._assert_published_error(result)
                self.assertIn("lectura posterior fallida", " ".join(result.errors))
                self.assertEqual(
                    list(self.output_path.parent.iterdir()), [self.output_path]
                )

    def test_published_digest_mismatch_reports_publication_with_error(self) -> None:
        original_digest = TemplateWriter._file_digest

        def mismatch_output_digest(writer, path):
            if path == self.output_path:
                return b"digest distinto"
            return original_digest(writer, path)

        with patch.object(TemplateWriter, "_file_digest", mismatch_output_digest):
            result = fill_template(self.template_path, self.weeks[:1], self.output_path)

        self._assert_published_error(result)
        self.assertIn("no coincide", " ".join(result.errors))

    def test_postpublication_stat_failure_preserves_published_result(self) -> None:
        original_stat = Path.stat
        original_replace = os.replace
        published = False

        def replace_then_mark(source, destination):
            nonlocal published
            original_replace(source, destination)
            published = True

        def fail_published_stat(path, *args, **kwargs):
            if published and path == self.output_path:
                raise PermissionError("consulta posterior fallida")
            return original_stat(path, *args, **kwargs)

        with (
            patch("meeting_generator.template_writer.os.replace", replace_then_mark),
            patch.object(Path, "stat", fail_published_stat),
        ):
            result = fill_template(self.template_path, self.weeks[:1], self.output_path)

        self._assert_published_error(result)

    def test_postpublication_resolve_failure_cannot_erase_publication(self) -> None:
        original_resolve = Path.resolve
        original_replace = os.replace
        published = False
        resolve_failed = False

        def replace_then_mark(source, destination):
            nonlocal published
            original_replace(source, destination)
            published = True

        def fail_published_resolve(path, *args, **kwargs):
            nonlocal resolve_failed
            if published and path == self.output_path:
                resolve_failed = True
                raise OSError("ruta posterior fallida")
            return original_resolve(path, *args, **kwargs)

        with (
            patch("meeting_generator.template_writer.os.replace", replace_then_mark),
            patch.object(Path, "resolve", fail_published_resolve),
        ):
            result = fill_template(self.template_path, self.weeks[:1], self.output_path)

        self.assertTrue(result.published)
        self.assertEqual(result.written_weeks, 1)
        self.assertEqual(result.omitted_weeks, 0)
        self.assertEqual(result.output_path, self.output_path)
        if resolve_failed:
            self._assert_published_error(result)
        else:
            self.assertTrue(result.succeeded, result.errors)

    def test_published_temporary_is_not_unlinked_after_replace(self) -> None:
        original_unlink = Path.unlink
        original_replace = os.replace
        published = False

        def replace_then_mark(source, destination):
            nonlocal published
            original_replace(source, destination)
            published = True

        def fail_temporary_cleanup(path, *args, **kwargs):
            if published and path.name.startswith(".published-"):
                raise PermissionError("limpieza posterior fallida")
            return original_unlink(path, *args, **kwargs)

        with (
            patch("meeting_generator.template_writer.os.replace", replace_then_mark),
            patch.object(
                Path, "unlink", autospec=True, side_effect=fail_temporary_cleanup
            ) as unlink,
        ):
            result = fill_template(self.template_path, self.weeks[:1], self.output_path)

        unlink.assert_not_called()
        self.assertTrue(result.published)
        self.assertTrue(result.succeeded, result.errors)
        self.assertEqual(result.written_weeks, 1)
        self.assertEqual(result.omitted_weeks, 0)
        self.assertEqual(result.output_path, self.output_path)
        self.assertEqual(list(self.output_path.parent.iterdir()), [self.output_path])

    def test_writer_logging_failure_after_save_preserves_published_result(self) -> None:
        def fail_completion_log(message, *args, **kwargs):
            if message == "Documento guardado y verificado.":
                raise OSError("log posterior fallido")

        with patch(
            "meeting_generator.template_writer.logger.info", fail_completion_log
        ):
            result = fill_template(self.template_path, self.weeks[:1], self.output_path)

        self._assert_published_error(result)
        self.assertIn("log posterior fallido", " ".join(result.errors))

    def test_reused_writer_does_not_report_previous_publication_on_new_failure(
        self,
    ) -> None:
        writer = TemplateWriter(self.template_path)
        successful = writer.fill(self.weeks[:1], self.output_path)
        self.assertTrue(successful.succeeded, successful.errors)
        original = self.output_path.read_bytes()

        with patch.object(writer, "_fill_document", side_effect=OSError("fallo nuevo")):
            result = writer.fill(self.weeks[:1], self.output_path)

        self._assert_unpublished(result)
        self.assertEqual(self.output_path.read_bytes(), original)

    def test_symlink_destination_reports_replaced_link_without_changing_target(
        self,
    ) -> None:
        target = self.output_path.with_name("previous.docx")
        target.write_bytes(b"destino original")
        self.output_path.symlink_to(target)

        result = fill_template(self.template_path, self.weeks[:1], self.output_path)

        self.assertTrue(result.succeeded, result.errors)
        self.assertTrue(result.published)
        self.assertEqual(result.output_path, self.output_path)
        self.assertFalse(self.output_path.is_symlink())
        self.assertEqual(target.read_bytes(), b"destino original")
        self.assertTrue(Document(self.output_path).tables)

    def test_generator_preserves_nine_published_weeks_after_output_read_failure(
        self,
    ) -> None:
        original_digest = TemplateWriter._file_digest

        def fail_output_digest(writer, path):
            if path == self.output_path:
                raise OSError("lectura posterior fallida")
            return original_digest(writer, path)

        with patch.object(TemplateWriter, "_file_digest", fail_output_digest):
            result = generate_document(
                self.source_path, self.template_path, self.output_path
            )

        self._assert_published_error(result, count=9)

    def test_generator_logging_failure_after_publication_returns_published_error(
        self,
    ) -> None:
        with patch(
            "meeting_generator.generator.logger.info",
            side_effect=OSError("log del resultado fallido"),
        ):
            result = generate_document(
                self.source_path, self.template_path, self.output_path
            )

        self._assert_published_error(result, count=9)
        self.assertIn("log del resultado fallido", " ".join(result.errors))


if __name__ == "__main__":
    unittest.main()
