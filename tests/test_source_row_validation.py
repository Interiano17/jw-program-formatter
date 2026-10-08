"""Rechazo de filas fuente ambiguas antes de publicar el programa."""

import copy
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from meeting_generator.generator import generate_document
from meeting_generator.parser import parse_document
from tests.source_fixtures import build_primary_source
from tests.template_fixture import build_template


class SourceRowValidationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        fixtures = tempfile.TemporaryDirectory()
        cls.addClassCleanup(fixtures.cleanup)
        directory = Path(fixtures.name)
        cls.source_path = build_primary_source(directory / "source.docx")
        cls.template_path = build_template(directory / "template.docx")
        cls.weeks = parse_document(cls.source_path)

    def test_ambiguous_rows_never_publish_or_replace_output(self) -> None:
        cases = (
            "omitted_number_cell",
            "omitted_number_with_grid_span",
            "omitted_participant_cell",
            "merged_assignment_columns",
            "number_merged_with_previous_row",
            "assignment_mixed_with_final_prayer",
            "conclusion_song_and_extra_assignment",
            "conclusion_inline_extra_assignment",
            "duration_without_number_or_title",
            "number_and_duration_without_title",
            "participant_without_number_or_content",
            "song_mixed_with_duration",
            "one_number_two_participant_groups",
            "one_number_two_assignments_shared_participant",
        )
        previous = b"programa anterior"
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            for case in cases:
                source_path = directory / f"{case}.docx"
                document = Document(self.source_path)
                row_number, row = next(
                    (number, row)
                    for number, row in enumerate(document.tables[1].rows, start=1)
                    if (
                        row.cells[1].text.startswith("PALABRAS DE CONCLUSIÓN")
                        if case.startswith("conclusion_")
                        else row.cells[0].text == "4"
                    )
                )
                if case.startswith("omitted_number"):
                    title_cell = row.cells[1]
                    row._tr.remove(row.cells[0]._tc)
                    grid_before = OxmlElement("w:gridBefore")
                    grid_before.set(qn("w:val"), "1")
                    row._tr.get_or_add_trPr().append(grid_before)
                    if case == "omitted_number_with_grid_span":
                        title_cell._tc.get_or_add_tcPr().get_or_add_gridSpan().val = 2
                elif case == "omitted_participant_cell":
                    row._tr.remove(row.cells[2]._tc)
                    grid_after = OxmlElement("w:gridAfter")
                    grid_after.set(qn("w:val"), "1")
                    row._tr.get_or_add_trPr().append(grid_after)
                elif case == "merged_assignment_columns":
                    row.cells[1].merge(row.cells[2])
                elif case == "number_merged_with_previous_row":
                    previous_row = document.tables[1].rows[row_number - 2]
                    previous_row.cells[0].merge(row.cells[0])
                elif case == "conclusion_song_and_extra_assignment":
                    row.cells[1].add_paragraph("CANCIÓN 22")
                    row.cells[1].add_paragraph("Práctica adicional (3 mins.)")
                elif case == "conclusion_inline_extra_assignment":
                    row.cells[1].text += " Práctica adicional"
                elif case.startswith("one_number_two"):
                    row.cells[1].add_paragraph("Práctica adicional (3 mins.)")
                    row.cells[2].text = "ALICIA EJEMPLO"
                    if case == "one_number_two_participant_groups":
                        row.cells[2].add_paragraph("BEATRIZ EJEMPLO")
                else:
                    if case != "number_and_duration_without_title":
                        row.cells[0].text = ""
                    row.cells[2].text = "ALICIA EJEMPLO"
                    row.cells[1].text = {
                        "assignment_mixed_with_final_prayer": (
                            "ORACIÓN FINAL\nPráctica sintética 1 (4 mins.)"
                        ),
                        "duration_without_number_or_title": "(4 mins.)",
                        "number_and_duration_without_title": "(4 mins.)",
                        "participant_without_number_or_content": "",
                        "song_mixed_with_duration": "CANCIÓN 49\n(4 mins.)",
                    }[case]
                document.save(source_path)
                source_bytes = source_path.read_bytes()
                template_bytes = self.template_path.read_bytes()

                for preexisting in (False, True):
                    with self.subTest(case=case, preexisting=preexisting):
                        output_path = directory / f"{case}-{preexisting}-output.docx"
                        if preexisting:
                            output_path.write_bytes(previous)
                        files_before = set(directory.iterdir())

                        with patch(
                            "meeting_generator.template_writer.os.replace",
                            wraps=os.replace,
                        ) as replace:
                            result = generate_document(
                                source_path, self.template_path, output_path
                            )

                        self.assertFalse(result.succeeded, result.errors)
                        self.assertFalse(result.published)
                        self.assertEqual(result.written_weeks, 0)
                        self.assertEqual(result.omitted_weeks, len(self.weeks))
                        self.assertIsNone(result.output_path)
                        self.assertTrue(
                            any(
                                f"semana 2, fila {row_number}" in error
                                for error in result.errors
                            ),
                            result.errors,
                        )
                        replace.assert_not_called()
                        self.assertEqual(set(directory.iterdir()), files_before)
                        self.assertEqual(source_path.read_bytes(), source_bytes)
                        self.assertEqual(
                            self.template_path.read_bytes(), template_bytes
                        )
                        if preexisting:
                            self.assertEqual(output_path.read_bytes(), previous)
                        else:
                            self.assertFalse(output_path.exists())

    def test_empty_rows_and_combined_heading_song_remain_valid(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source_path = Path(temporary) / "empty-rows.docx"
            document = Document(self.source_path)
            table = document.tables[0]
            table.rows[0].cells[0].merge(table.rows[0].cells[2])
            ministry_row = next(row for row in table.rows if row.cells[0].text == "4")
            for cell_count in range(4):
                empty_row = table.add_row()
                for cell in empty_row.cells[cell_count:]:
                    empty_row._tr.remove(cell._tc)
                if cell_count < 3:
                    omitted_grid = OxmlElement(
                        "w:gridBefore" if cell_count == 2 else "w:gridAfter"
                    )
                    omitted_grid.set(qn("w:val"), str(3 - cell_count))
                    empty_row._tr.get_or_add_trPr().append(omitted_grid)
                ministry_row._tr.addprevious(empty_row._tr)

            heading_index, heading = next(
                (index, row)
                for index, row in enumerate(table.rows)
                if len(row.cells) > 1 and row.cells[1].text == "NUESTRA VIDA CRISTIANA"
            )
            song = table.rows[heading_index + 1]
            heading.cells[1].add_paragraph(song.cells[1].text)
            table._tbl.remove(song._tr)
            conclusion = next(
                row
                for row in table.rows
                if len(row.cells) > 1
                and row.cells[1].text.startswith("PALABRAS DE CONCLUSIÓN")
            )
            closing_song = next(
                row
                for row in table.rows
                if len(row.cells) > 1
                and row.cells[1].text == f"CANCIÓN {self.weeks[0].closing_song}"
            )
            conclusion.cells[1].add_paragraph(closing_song.cells[1].text)
            table._tbl.remove(closing_song._tr)
            document.save(source_path)

            expected = copy.deepcopy(self.weeks)
            expected[0].christian_life[-1].raw_lines.append(
                f"CANCIÓN {self.weeks[0].closing_song}"
            )
            self.assertEqual(parse_document(source_path), expected)


if __name__ == "__main__":
    unittest.main()
