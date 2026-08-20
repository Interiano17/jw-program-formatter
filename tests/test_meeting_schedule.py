import tempfile
import unittest
from pathlib import Path

from docx import Document

from meeting_generator.parser import parse_document
from meeting_generator.template_writer import fill_template
from meeting_generator.utils import normalize_person_name, normalize_sentence_case


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_DOCUMENT = PROJECT_ROOT / "JULIO-AGOSTO-REVISADO 2.docx"
SECOND_SOURCE_DOCUMENT = PROJECT_ROOT / "SEPTIEMBRE-NOVIEMBRE.docx"
TEMPLATE_DOCUMENT = PROJECT_ROOT / "S-140_S.docx"


class ParserRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.weeks = parse_document(SOURCE_DOCUMENT)

    def test_numbered_section_song_is_not_an_assignment(self) -> None:
        week = self.weeks[5]

        normal_assignments = [
            assignment
            for assignment in week.christian_life
            if not assignment.is_bible_study and not assignment.is_conclusion
        ]

        self.assertEqual([assignment.number for assignment in normal_assignments], [8])
        self.assertEqual(normal_assignments[0].duration_minutes, 15)

    def test_combined_row_is_split_and_missing_duration_is_inferred(self) -> None:
        week = self.weeks[7]
        normal_assignments = [
            assignment
            for assignment in week.christian_life
            if not assignment.is_bible_study and not assignment.is_conclusion
        ]

        self.assertEqual([assignment.number for assignment in normal_assignments], [7, 8])
        self.assertEqual(
            [assignment.duration_minutes for assignment in normal_assignments],
            [10, 5],
        )
        self.assertEqual(normal_assignments[1].duration_text, "5 mins.")

    def test_intermediate_and_closing_songs_are_distinct(self) -> None:
        first_week = self.weeks[0]

        self.assertEqual(first_week.intermediate_song, "49")
        self.assertEqual(first_week.closing_song, "61")


class ScheduleIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temp_directory = tempfile.TemporaryDirectory()
        output_path = Path(cls.temp_directory.name) / "generated.docx"
        weeks = parse_document(SOURCE_DOCUMENT)
        fill_template(TEMPLATE_DOCUMENT, weeks, output_path)
        cls.document = Document(output_path)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temp_directory.cleanup()

    def test_first_week_uses_expected_start_times(self) -> None:
        first_copy = self.document.tables[0].rows[:27]
        times = [
            row.cells[0].text.strip()
            for row in first_copy
            if row.cells and ":" in row.cells[0].text.strip()
        ]

        self.assertEqual(
            times,
            [
                "7:00",
                "7:05",
                "7:06",
                "7:16",
                "7:26",
                "7:30",
                "7:34",
                "7:39",
                "7:45",
                "7:49",
                "8:04",
                "8:34",
                "8:37",
            ],
        )

    def test_ministry_extra_minute_does_not_change_displayed_duration(self) -> None:
        rows = self.document.tables[0].rows

        self.assertEqual(rows[14].cells[0].text.strip(), "7:30")
        self.assertIn("(3 mins.)", rows[14].cells[1].text)
        self.assertEqual(rows[15].cells[0].text.strip(), "7:34")

    def test_weekly_reading_replaces_template_label(self) -> None:
        header = self.document.tables[0].rows[2].cells[0].text

        self.assertEqual(header, "6 al 12 julio | Jeremías 13-15")
        self.assertNotIn("LECTURA SEMANAL DE LA BIBLIA", header)

    def test_two_participants_use_one_slash(self) -> None:
        participant_text = self.document.tables[0].rows[14].cells[-1].text

        self.assertEqual(participant_text, "Lina Núñez / Nevelin Andino")
        self.assertNotIn("//", participant_text)

    def test_student_helper_label_is_preserved(self) -> None:
        row_texts = [
            cell.text.strip()
            for cell in self.document.tables[0].rows[14].cells
        ]

        self.assertIn("Estudiante/Ayudante:", row_texts)

    def test_congregation_name_is_written_in_headers(self) -> None:
        document_text = "\n".join(
            cell.text
            for table in self.document.tables
            for row in table.rows
            for cell in row.cells
        )

        self.assertIn("Congregación Loarque", document_text)
        self.assertNotIn("[NOMBRE DE LA CONGREGACIÓN]", document_text)

    def test_combined_christian_life_row_has_correct_times(self) -> None:
        rows = self.document.tables[2].rows

        self.assertEqual(rows[46].cells[0].text.strip(), "7:49")
        self.assertIn("(10 mins.)", rows[46].cells[2].text)
        self.assertEqual(rows[47].cells[0].text.strip(), "7:59")
        self.assertIn("(5 mins.)", rows[47].cells[2].text)
        self.assertEqual(rows[48].cells[0].text.strip(), "8:04")

    def test_generated_document_has_no_zero_time_markers(self) -> None:
        for table in self.document.tables:
            for row in table.rows:
                for cell in row.cells:
                    self.assertNotIn("0:00", cell.text)


class SecondDocumentRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.weeks = parse_document(SECOND_SOURCE_DOCUMENT)
        cls.temp_directory = tempfile.TemporaryDirectory()
        output_path = Path(cls.temp_directory.name) / "generated.docx"
        fill_template(TEMPLATE_DOCUMENT, cls.weeks, output_path)
        cls.document = Document(output_path)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temp_directory.cleanup()

    def test_header_and_bible_study_format_variations_are_accepted(self) -> None:
        third_week = self.weeks[2]

        self.assertEqual(third_week.president, "Jezer Zúniga")
        self.assertEqual(third_week.opening_song, "74")
        self.assertEqual(
            [
                assignment.number
                for assignment in third_week.christian_life
                if assignment.is_bible_study
            ],
            [8],
        )

    def test_unnumbered_christian_life_assignments_receive_number(self) -> None:
        for week_index in (4, 5):
            normal_assignments = [
                assignment
                for assignment in self.weeks[week_index].christian_life
                if not assignment.is_bible_study and not assignment.is_conclusion
            ]
            self.assertEqual(
                [(assignment.number, assignment.duration_minutes)
                 for assignment in normal_assignments],
                [(8, 15)],
            )

    def test_ministry_discourse_without_duration_uses_five_minutes(self) -> None:
        discourse = self.weeks[6].ministry[-1]

        self.assertEqual(discourse.number, 7)
        self.assertEqual(discourse.duration_minutes, 5)
        self.assertEqual(discourse.duration_text, "5 mins.")

    def test_variable_ministry_rows_generate_expected_times(self) -> None:
        first_copy = self.document.tables[0].rows[:27]
        times = [
            row.cells[0].text.strip()
            for row in first_copy
            if row.cells and ":" in row.cells[0].text.strip()
        ]

        self.assertEqual(
            times,
            [
                "7:00",
                "7:05",
                "7:06",
                "7:16",
                "7:26",
                "7:30",
                "7:34",
                "7:39",
                "7:43",
                "7:58",
                "8:28",
                "8:31",
            ],
        )

    def test_second_generated_document_has_no_zero_time_markers(self) -> None:
        for table in self.document.tables:
            for row in table.rows:
                for cell in row.cells:
                    self.assertNotIn("0:00", cell.text)

    def test_uppercase_source_names_are_normalized(self) -> None:
        first_week = self.weeks[0]

        self.assertEqual(first_week.president, "Marvin Zúniga")
        self.assertEqual(
            first_week.bible_treasures[0].first_participant_name(),
            "Jezer Zúniga",
        )

    def test_second_document_includes_weekly_reading(self) -> None:
        header = self.document.tables[0].rows[2].cells[0].text

        self.assertEqual(
            header,
            "7 al 13 de septiembre | Jeremías 32-33",
        )


class NameNormalizationTests(unittest.TestCase):
    def test_only_fully_uppercase_names_are_changed(self) -> None:
        self.assertEqual(normalize_person_name("JUAN PÉREZ"), "Juan Pérez")
        self.assertEqual(normalize_person_name("Juan McDONALD"), "Juan McDONALD")

    def test_dates_and_readings_use_sentence_case(self) -> None:
        self.assertEqual(
            normalize_sentence_case("7 AL 13 DE SEPTIEMBRE"),
            "7 al 13 de septiembre",
        )
        self.assertEqual(
            normalize_sentence_case("JEREMÍAS 32-33"),
            "Jeremías 32-33",
        )


if __name__ == "__main__":
    unittest.main()
