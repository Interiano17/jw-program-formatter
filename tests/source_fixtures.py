"""Constructores de documentos fuente sintéticos para las pruebas.

Los archivos generados contienen solo la estructura que consume
``ProgramParser`` y nombres ficticios. No dependen de documentos operativos ni
de datos personales externos al repositorio.
"""

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table, _Cell


@dataclass(frozen=True)
class AssignmentFixture:
    """Datos mínimos para crear una fila analizable."""

    number: int | None
    title: str
    duration_minutes: int | None
    participants: str


@dataclass(frozen=True)
class WeekFixture:
    """Configuración de una semana sintética."""

    date: str
    reading: str
    president: str
    opening_song: int
    intermediate_song: int
    closing_song: int
    ministry: tuple[AssignmentFixture, ...]
    christian_life: tuple[AssignmentFixture, ...]
    bible_study_number: int
    conclusion_number: int
    treasure_participant: str
    numbered_section_song: int | None = None
    combined_christian_life: bool = False


_TWO_PARTICIPANTS = "ALICIA EJEMPLO // BEATRIZ EJEMPLO"
_ONE_PARTICIPANT = "CARLOS EJEMPLO"
_STUDY_PARTICIPANTS = "DIANA EJEMPLO // ERNESTO EJEMPLO"


def build_primary_source(path: Path) -> Path:
    """Genera la fuente primaria con las variantes de parser bajo regresión."""
    weeks = (
        _standard_week(
            "6 AL 12 JULIO",
            "JEREMÍAS 13-15",
            opening_song=123,
            intermediate_song=49,
            closing_song=61,
            ministry_durations=(3, 4, 5),
        ),
        _standard_week(
            "13 AL 19 JULIO",
            "ISAÍAS 1-3",
            opening_song=34,
            intermediate_song=54,
            closing_song=22,
        ),
        _standard_week(
            "20 AL 26 JULIO",
            "OSEAS 1-3",
            opening_song=44,
            intermediate_song=38,
            closing_song=153,
        ),
        _standard_week(
            "27 DE JULIO AL 2 DE AGOSTO",
            "1 CORINTIOS 1-2",
            opening_song=73,
            intermediate_song=57,
            closing_song=55,
        ),
        _standard_week(
            "3 AL 9 AGOSTO",
            "JEREMÍAS 22-23",
            opening_song=40,
            intermediate_song=103,
            closing_song=60,
        ),
        _standard_week(
            "10 AL 16 DE AGOSTO",
            "JEREMÍAS 24-25",
            opening_song=124,
            intermediate_song=65,
            closing_song=137,
            normal_number=8,
            bible_study_number=9,
            conclusion_number=10,
            numbered_section_song=7,
        ),
        _standard_week(
            "17 AL 23 DE AGOSTO",
            "JEREMÍAS 26-28",
            opening_song=77,
            intermediate_song=16,
            closing_song=71,
        ),
        _standard_week(
            "24 AL 30 DE AGOSTO",
            "JEREMÍAS 29-30",
            opening_song=12,
            intermediate_song=3,
            closing_song=156,
            christian_life_durations=(10, 5),
            bible_study_number=9,
            conclusion_number=10,
            combined_christian_life=True,
        ),
        _standard_week(
            "31 DE AGOSTO AL 6 DE SEPTIEMBRE",
            "JEREMÍAS 31",
            opening_song=27,
            intermediate_song=67,
            closing_song=132,
        ),
    )
    return _build_source(path, weeks)


def build_secondary_source(path: Path) -> Path:
    """Genera la fuente secundaria y sus variaciones de formato."""
    weeks = (
        _standard_week(
            "7 AL 13 DE SEPTIEMBRE",
            "JEREMÍAS 32-33",
            president="MARIO PRUEBA",
            opening_song=1,
            intermediate_song=128,
            closing_song=143,
            ministry_durations=(3, 4),
            normal_number=7,
            bible_study_number=8,
            conclusion_number=9,
            treasure_participant="ELENA PRUEBA",
        ),
        _standard_week(
            "14 AL 20 DE SEPTIEMBRE",
            "JEREMÍAS 34-35",
            opening_song=161,
            intermediate_song=121,
            closing_song=28,
            ministry_durations=(2, 2, 3, 4),
            christian_life_durations=(6, 9),
            normal_number=8,
            bible_study_number=10,
            conclusion_number=11,
        ),
        _standard_week(
            "21 AL 27 DE SEPTIEMBRE",
            "JEREMÍAS 36-37",
            president="Celia Ejemplo",
            opening_song=74,
            intermediate_song=142,
            closing_song=134,
            ministry_durations=(3, 4, 6),
        ),
        _standard_week(
            "28 DE SEPTIEMBRE AL 4 DE OCTUBRE",
            "JEREMÍAS 38-39",
            opening_song=102,
            intermediate_song=90,
            closing_song=56,
            ministry_durations=(3, 4, 6),
        ),
        _standard_week(
            "5 AL 11 DE OCTUBRE",
            "JEREMÍAS 40-41",
            opening_song=33,
            intermediate_song=17,
            closing_song=38,
            ministry_durations=(2, 2, 4, 3),
            normal_number=None,
            bible_study_number=9,
            conclusion_number=10,
        ),
        _standard_week(
            "12 AL 18 DE OCTUBRE",
            "JEREMÍAS 42-44",
            opening_song=103,
            intermediate_song=47,
            closing_song=129,
            ministry_durations=(3, 4, 6),
            normal_number=None,
            bible_study_number=9,
            conclusion_number=10,
            numbered_section_song=7,
        ),
        _standard_week(
            "19 AL 25 DE OCTUBRE",
            "JEREMÍAS 45-46",
            opening_song=21,
            intermediate_song=117,
            closing_song=87,
            ministry_durations=(3, 2, 2, None),
            normal_number=8,
            bible_study_number=9,
            conclusion_number=10,
        ),
        _standard_week(
            "26 DE OCTUBRE AL 1 DE NOVIEMBRE",
            "JEREMÍAS 47-48",
            opening_song=125,
            intermediate_song=158,
            closing_song=54,
        ),
    )
    return _build_source(path, weeks)


def build_source_with_auxiliary_table(path: Path) -> Path:
    """Genera dos semanas separadas por una tabla ajena al programa."""
    weeks = (
        _standard_week(
            "1 AL 7 ENERO",
            "ISAÍAS 1-3",
            opening_song=1,
            intermediate_song=2,
            closing_song=3,
        ),
        _standard_week(
            "8 AL 14 ENERO",
            "OSEAS 1-3",
            opening_song=4,
            intermediate_song=5,
            closing_song=6,
        ),
    )
    return _build_source(path, weeks, add_auxiliary_after_first=True)


def _standard_week(
    date: str,
    reading: str,
    *,
    president: str = "ANA EJEMPLO",
    opening_song: int,
    intermediate_song: int,
    closing_song: int,
    ministry_durations: tuple[int | None, ...] = (4, 4, 4),
    christian_life_durations: tuple[int, ...] = (15,),
    normal_number: int | None = 7,
    bible_study_number: int = 8,
    conclusion_number: int = 9,
    numbered_section_song: int | None = None,
    combined_christian_life: bool = False,
    treasure_participant: str = _ONE_PARTICIPANT,
) -> WeekFixture:
    ministry = tuple(
        AssignmentFixture(
            number=4 + index,
            title=(
                "DISCURSO sintético"
                if duration is None
                else f"Práctica sintética {index + 1}"
            ),
            duration_minutes=duration,
            participants=(
                _ONE_PARTICIPANT if duration is None else _TWO_PARTICIPANTS
            ),
        )
        for index, duration in enumerate(ministry_durations)
    )

    christian_life = tuple(
        AssignmentFixture(
            number=(normal_number + index if normal_number is not None else None),
            title=f"Tema sintético de vida cristiana {index + 1}",
            duration_minutes=duration,
            participants=_ONE_PARTICIPANT,
        )
        for index, duration in enumerate(christian_life_durations)
    )

    return WeekFixture(
        date=date,
        reading=reading,
        president=president,
        opening_song=opening_song,
        intermediate_song=intermediate_song,
        closing_song=closing_song,
        ministry=ministry,
        christian_life=christian_life,
        bible_study_number=bible_study_number,
        conclusion_number=conclusion_number,
        treasure_participant=treasure_participant,
        numbered_section_song=numbered_section_song,
        combined_christian_life=combined_christian_life,
    )


def _build_source(
    path: Path,
    weeks: tuple[WeekFixture, ...],
    *,
    add_auxiliary_after_first: bool = False,
) -> Path:
    document = Document()
    document.core_properties.title = "Fuente sintética para pruebas"
    document.core_properties.author = ""

    for index, week in enumerate(weeks):
        _add_week_header(document, week)
        _add_week_table(document, week)
        if add_auxiliary_after_first and index == 0:
            document.add_table(rows=1, cols=1).cell(0, 0).text = "Nota auxiliar"

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def _add_week_header(document: DocxDocument, week: WeekFixture) -> None:
    document.add_paragraph(f"SEMANA DEL {week.date}. {week.reading}")
    document.add_paragraph(
        f"PRESIDENTE: {week.president}. CANCIÓN {week.opening_song}"
    )


def _add_week_table(document: DocxDocument, week: WeekFixture) -> None:
    table = document.add_table(rows=0, cols=3)
    _add_special_row(table, "TESOROS DE LA BIBLIA")
    for number, duration, participant in (
        (1, 10, week.treasure_participant),
        (2, 10, _ONE_PARTICIPANT),
        (3, 4, _ONE_PARTICIPANT),
    ):
        _add_assignment_row(
            table,
            AssignmentFixture(
                number,
                f"Tema sintético de tesoros {number}",
                duration,
                participant,
            ),
        )

    _add_special_row(table, "SEAMOS MEJORES MAESTROS")
    for assignment in week.ministry:
        _add_assignment_row(table, assignment)

    if week.numbered_section_song is not None:
        cells = table.add_row().cells
        cells[0].text = str(week.numbered_section_song)
        _set_cell_paragraphs(
            cells[1],
            ("NUESTRA VIDA CRISTIANA", f"CANCIÓN {week.intermediate_song}"),
        )
    else:
        _add_special_row(table, "NUESTRA VIDA CRISTIANA")
        _add_special_row(table, f"CANCIÓN {week.intermediate_song}")

    if week.combined_christian_life:
        _add_combined_christian_life_row(table, week.christian_life)
    else:
        for assignment in week.christian_life:
            _add_assignment_row(table, assignment)

    _add_assignment_row(
        table,
        AssignmentFixture(
            week.bible_study_number,
            "Estudio bíblico de la congregación",
            30,
            _STUDY_PARTICIPANTS,
        ),
    )
    _add_assignment_row(
        table,
        AssignmentFixture(
            week.conclusion_number,
            "PALABRAS DE CONCLUSIÓN",
            3,
            _ONE_PARTICIPANT,
        ),
    )
    _add_special_row(table, f"CANCIÓN {week.closing_song}")
    _add_special_row(table, "ORACIÓN FINAL", _ONE_PARTICIPANT)


def _add_assignment_row(table: Table, assignment: AssignmentFixture) -> None:
    cells = table.add_row().cells
    cells[0].text = "" if assignment.number is None else str(assignment.number)
    suffix = (
        f" ({assignment.duration_minutes} mins.)"
        if assignment.duration_minutes is not None
        else ""
    )
    cells[1].text = assignment.title + suffix
    cells[2].text = assignment.participants


def _add_combined_christian_life_row(
    table: Table,
    assignments: tuple[AssignmentFixture, ...],
) -> None:
    if len(assignments) != 2:
        raise ValueError("la fila combinada sintética requiere dos asignaciones")

    first, second = assignments
    _set_cell_paragraphs(
        table.add_row().cells[0],
        (str(first.number), str(second.number)),
    )
    row = table.rows[-1].cells
    _set_cell_paragraphs(
        row[1],
        (
            f"{first.title} ({first.duration_minutes} mins.)",
            second.title,
        ),
    )
    _set_cell_paragraphs(
        row[2],
        (first.participants, second.participants),
    )


def _add_special_row(table: Table, text: str, participant: str = "") -> None:
    cells = table.add_row().cells
    cells[1].text = text
    cells[2].text = participant


def _set_cell_paragraphs(cell: _Cell, values: tuple[str, ...]) -> None:
    cell.text = values[0]
    for value in values[1:]:
        cell.add_paragraph(value)
