"""
Módulo parser para el Meeting Generator.

Analiza el documento fuente (.docx) del programa de reunión
y extrae la información estructurada de cada semana.

Estructura real del documento:
- Encabezados de semana (SEMANA DEL..., PRESIDENTE:...) en párrafos.
- Contenido de cada semana (puntos) en TABLAS de 3 columnas.
- Las secciones se detectan por sus encabezados, no por rango de números.

El resultado solo se devuelve si todas las semanas cumplen el contrato.
"""

import logging
from pathlib import Path
from typing import Optional

from docx import Document
from docx.document import Document as DocxDocument

from .config import (
    CHRISTIAN_LIFE_PRE_STUDY_MINUTES,
    DEFAULT_DURATIONS,
    RE_BIBLE_STUDY,
    RE_CONCLUSION_WORDS,
    RE_DATE,
    RE_FINAL_PRAYER,
    RE_PRESIDENT_SONG,
    RE_SONG,
    RE_WEEK_HEADER,
    RE_WEEKLY_READING,
)
from .models import Assignment, MeetingWeek, Participant
from .utils import (
    clean_text,
    extract_number,
    normalize_person_name,
    normalize_sentence_case,
    remove_name_suffix,
    split_names,
)
from .validation import GenerationValidationError

logger = logging.getLogger(__name__)


class ProgramParser:
    """Parser del documento de programa de reunión."""

    def __init__(self, file_path: Path):
        self.file_path: Path = file_path
        self.document: Optional[DocxDocument] = None
        self.weeks: list[MeetingWeek] = []
        self.detected_week_count = 0

    def parse(self) -> list[MeetingWeek]:
        """Analiza y valida el documento sin devolver resultados parciales."""
        logger.info("Iniciando análisis: %s", self.file_path)
        self.document = Document(str(self.file_path))
        self.weeks = []
        self.detected_week_count = 0
        issues: list[str] = []

        week_headers = self._extract_week_headers_from_paragraphs()
        logger.info("Encabezados encontrados: %d semanas", len(week_headers))

        tables = self.document.tables
        program_tables = [table for table in tables if self._is_program_table(table)]
        logger.info(
            "Tablas encontradas: %d; tablas de programa: %d",
            len(tables),
            len(program_tables),
        )
        self.detected_week_count = max(len(week_headers), len(program_tables))
        if not program_tables:
            issues.append("el documento no contiene tablas de semanas")
        if len(week_headers) != len(program_tables):
            issues.append(
                "la cantidad de encabezados "
                f"({len(week_headers)}) no coincide con la de tablas de programa "
                f"({len(program_tables)})"
            )

        for week_index, table in enumerate(program_tables):
            try:
                header_data = (
                    week_headers[week_index]
                    if week_index < len(week_headers)
                    else {}
                )
                meeting_week = self._parse_week_from_table(
                    table,
                    week_index + 1,
                    header_data,
                )

                logger.info(
                    "Sem %d: %s | Pres=%s | T=%d M=%d CV=%d",
                    week_index + 1,
                    meeting_week.date,
                    meeting_week.president,
                    len(meeting_week.bible_treasures),
                    len(meeting_week.ministry),
                    len(meeting_week.christian_life),
                )
                self.weeks.append(meeting_week)

            except Exception as exc:
                logger.error("Error tabla %d: %s", week_index + 1, exc, exc_info=True)
                issues.append(
                    f"tabla {week_index + 1}: no se pudo analizar "
                    f"({type(exc).__name__}: {exc})"
                )

        for week in self.weeks:
            if not week.is_valid():
                issues.extend(
                    f"semana {week.week_index}: {issue}"
                    for issue in week.validation_errors()
                )

        if issues:
            raise GenerationValidationError("Documento fuente inválido", issues)

        logger.info("Análisis completado. %d semanas extraídas.", len(self.weeks))
        return self.weeks

    def _is_program_table(self, table) -> bool:
        """Distingue una tabla de semana de una tabla auxiliar."""
        sections: set[str] = set()
        for row in table.rows:
            if len(row.cells) > 1:
                section = self._detect_section_from_lines(
                    self._split_cell_lines(row.cells[1].text)
                )
                if section:
                    sections.add(section)
        return sections == {"treasures", "ministry", "christian_life"}

    # ------------------------------------------------------------------
    # Encabezados desde párrafos (sin cambios)
    # ------------------------------------------------------------------

    def _extract_week_headers_from_paragraphs(self) -> list[dict]:
        """Extrae encabezados de semana desde los párrafos."""
        if self.document is None:
            return []

        paragraphs = [p.text for p in self.document.paragraphs]
        headers: list[dict] = []
        current_header: dict = {}

        for para_text in paragraphs:
            stripped = para_text.strip()
            if not stripped:
                continue

            if RE_WEEK_HEADER.match(stripped):
                if current_header:
                    headers.append(current_header)
                current_header = {"raw": stripped}
                date, weekly_reading = self._extract_week_header_fields(stripped)
                current_header["date"] = date
                if weekly_reading:
                    current_header["weekly_reading"] = weekly_reading
                continue

            pres_match = RE_PRESIDENT_SONG.search(stripped)
            if pres_match and "president" not in current_header:
                current_header["president"] = normalize_person_name(
                    pres_match.group(1)
                )
                current_header["opening_song"] = clean_text(pres_match.group(2))
                continue

            if "weekly_reading" not in current_header and current_header:
                read_match = RE_WEEKLY_READING.search(stripped)
                if read_match:
                    current_header["weekly_reading"] = self._normalize_weekly_reading(
                        read_match.group(0)
                    )

        if current_header:
            headers.append(current_header)
        return headers

    @staticmethod
    def _extract_week_header_fields(header: str) -> tuple[str, str]:
        """Separa fecha y lectura usando la posición real de la lectura."""
        cleaned_header = clean_text(header)
        read_match = RE_WEEKLY_READING.search(cleaned_header)
        date_end = read_match.start() if read_match else len(cleaned_header)
        date_match = RE_DATE.search(cleaned_header[:date_end])
        date_text = date_match.group(1) if date_match else ""
        date = normalize_sentence_case(date_text.rstrip(" \t.,;:-–—"))
        weekly_reading = (
            ProgramParser._normalize_weekly_reading(read_match.group(0))
            if read_match
            else ""
        )
        return date, weekly_reading

    @staticmethod
    def _normalize_weekly_reading(reading: str) -> str:
        """Normaliza la lectura y conserva la inicial tras un ordinal."""
        normalized = normalize_sentence_case(reading)
        for index, character in enumerate(normalized):
            if character.isalpha():
                return (
                    normalized[:index]
                    + character.upper()
                    + normalized[index + 1 :]
                )
        return normalized

    # ------------------------------------------------------------------
    # Parseo de semana desde tabla
    # ------------------------------------------------------------------

    def _parse_week_from_table(
        self, table, week_number: int, header_data: dict
    ) -> MeetingWeek:
        """
        Parsea una semana desde una tabla del documento.

        Estrategia dinámica:
        - Mantiene la sección actual ("treasures", "ministry", "christian_life").
        - Cambia de sección al encontrar encabezados.
        - Cada punto se convierte en un Assignment con su número original.
        """
        week = MeetingWeek(week_index=week_number)

        week.raw_header = header_data.get("raw", "")
        week.date = header_data.get("date", "")
        week.weekly_reading = header_data.get("weekly_reading", "")
        week.president = header_data.get("president", "")
        week.opening_song = header_data.get("opening_song", "")
        week.opening_prayer = week.president

        current_section: str = "treasures"

        for row in table.rows:
            cells = row.cells
            if len(cells) < 3:
                continue

            point_numbers = self._extract_point_numbers(cells[0])
            point_number = point_numbers[0] if len(point_numbers) == 1 else None
            cell1_text = cells[1].text
            lines = self._split_cell_lines(cell1_text)
            person_text = clean_text(cells[2].text)

            # --- Detectar cambio de sección por encabezado ---
            new_section = self._detect_section_from_lines(lines)
            if new_section:
                current_section = new_section
                if point_number is None:
                    # Fila sin número, solo encabezado
                    self._process_special_row(week, lines, person_text)
                    continue

            # --- Detectar palabras de conclusión ---
            if self._lines_contain_conclusion(lines):
                assignment = self._create_assignment(
                    point_number or 0, lines, person_text
                )
                assignment.is_conclusion = True
                week.christian_life.append(assignment)
                continue

            # --- Detectar oración final ---
            if self._lines_contain_final_prayer(lines):
                if person_text:
                    week.closing_prayer = normalize_person_name(
                        remove_name_suffix(person_text)
                    )
                continue

            # Algunas tablas agrupan varias asignaciones en una misma fila.
            # En ese caso los números, títulos y participantes se conservan en
            # párrafos separados y deben procesarse individualmente.
            if len(point_numbers) > 1:
                assignments = self._create_assignments_from_combined_row(
                    point_numbers,
                    cells[1],
                    cells[2],
                )
                self._append_assignments(week, current_section, assignments)
                continue

            # --- Crear assignment genérico ---
            if point_number is None:
                candidate = self._create_assignment(0, lines, person_text)
                if (
                    current_section == "christian_life"
                    and candidate.title.strip()
                    and (
                        candidate.duration_minutes > 0
                        or candidate.has_participants()
                    )
                ):
                    self._append_assignments(
                        week,
                        current_section,
                        [candidate],
                    )
                    continue
                self._process_special_row(week, lines, person_text)
                continue

            assignment = self._create_assignment(point_number, lines, person_text)

            # Hay documentos donde el encabezado de Vida Cristiana y su canción
            # aparecen en una fila que conserva por error un número de punto.
            if not assignment.title.strip():
                self._process_special_row(week, lines, person_text)
                continue

            # Asignar a la sección actual
            self._append_assignments(week, current_section, [assignment])

        # Extraer cierre (canción final, oración final)
        self._extract_closing_from_table(week, table)
        self._assign_missing_christian_life_numbers(week)
        self._infer_missing_ministry_durations(week)
        self._infer_missing_christian_life_duration(week)

        return week

    def _extract_point_numbers(self, cell) -> list[int]:
        """Extrae por separado los números escritos en párrafos de una celda."""
        numbers: list[int] = []
        for paragraph in cell.paragraphs:
            number = self._clean_point_number(clean_text(paragraph.text))
            if number is not None:
                numbers.append(number)
        return numbers

    def _create_assignments_from_combined_row(
        self,
        point_numbers: list[int],
        content_cell,
        participant_cell,
    ) -> list[Assignment]:
        """Separa una fila que contiene más de una asignación."""
        content_lines = [
            clean_text(paragraph.text)
            for paragraph in content_cell.paragraphs
            if clean_text(paragraph.text)
        ]
        participant_lines = [
            clean_text(paragraph.text)
            for paragraph in participant_cell.paragraphs
            if clean_text(paragraph.text)
        ]

        if len(content_lines) != len(point_numbers):
            logger.warning(
                "Fila combinada ambigua: números=%s, títulos=%s",
                point_numbers,
                content_lines,
            )

        assignments: list[Assignment] = []
        for index, number in enumerate(point_numbers):
            title_line = content_lines[index] if index < len(content_lines) else ""
            participant_text = (
                participant_lines[index] if index < len(participant_lines) else ""
            )
            assignments.append(
                self._create_assignment(number, [title_line], participant_text)
            )
        return assignments

    def _append_assignments(
        self,
        week: MeetingWeek,
        section: str,
        assignments: list[Assignment],
    ) -> None:
        """Agrega asignaciones a la sección activa."""
        if section == "treasures":
            week.bible_treasures.extend(assignments)
        elif section == "ministry":
            week.ministry.extend(assignments)
        elif section == "christian_life":
            week.christian_life.extend(assignments)

    def _infer_missing_christian_life_duration(self, week: MeetingWeek) -> None:
        """Completa una única duración omitida antes del estudio bíblico."""
        assignments = [
            assignment
            for assignment in week.christian_life
            if not assignment.is_bible_study
            and not assignment.is_conclusion
            and assignment.title.strip()
        ]
        missing = [
            assignment for assignment in assignments
            if assignment.duration_minutes <= 0
        ]
        known_minutes = sum(
            assignment.duration_minutes
            for assignment in assignments
            if assignment.duration_minutes > 0
        )
        remaining = CHRISTIAN_LIFE_PRE_STUDY_MINUTES - known_minutes

        if len(missing) == 1 and remaining > 0:
            missing[0].duration_minutes = remaining
            missing[0].duration_text = f"{remaining} mins."
            logger.info(
                "Semana %d, punto %d: duración inferida de %d minutos.",
                week.week_index,
                missing[0].number,
                remaining,
            )
        elif missing:
            logger.warning(
                "Semana %d: no se pudo inferir con seguridad la duración de %s.",
                week.week_index,
                [assignment.number for assignment in missing],
            )

    def _assign_missing_christian_life_numbers(self, week: MeetingWeek) -> None:
        """Asigna números omitidos usando la posición anterior al estudio."""
        unnumbered = [
            assignment
            for assignment in week.christian_life
            if assignment.number <= 0
            and not assignment.is_bible_study
            and not assignment.is_conclusion
        ]
        bible_study = next(
            (
                assignment
                for assignment in week.christian_life
                if assignment.is_bible_study and assignment.number > 0
            ),
            None,
        )
        if not unnumbered or bible_study is None:
            return

        used_numbers = {
            assignment.number
            for assignment in week.all_assignments()
            if assignment.number > 0
        }
        available_numbers = [
            number
            for number in range(1, bible_study.number)
            if number not in used_numbers
        ]
        inferred_numbers = available_numbers[-len(unnumbered):]
        if len(inferred_numbers) != len(unnumbered):
            logger.warning(
                "Semana %d: no se pudieron inferir números para %d asignaciones.",
                week.week_index,
                len(unnumbered),
            )
            return

        for assignment, number in zip(unnumbered, inferred_numbers):
            assignment.number = number
            logger.info(
                "Semana %d: número %d inferido para '%s'.",
                week.week_index,
                number,
                assignment.title,
            )

    def _infer_missing_ministry_durations(self, week: MeetingWeek) -> None:
        """Aplica la duración estándar a discursos sin tiempo explícito."""
        for assignment in week.ministry:
            if (
                assignment.duration_minutes <= 0
                and assignment.title.upper().startswith("DISCURSO")
            ):
                duration = DEFAULT_DURATIONS["ministerio_discurso"]
                assignment.duration_minutes = duration
                assignment.duration_text = f"{duration} mins."
                logger.info(
                    "Semana %d, punto %d: duración de discurso inferida en %d minutos.",
                    week.week_index,
                    assignment.number,
                    duration,
                )

    # ------------------------------------------------------------------
    # Detección de secciones por encabezado
    # ------------------------------------------------------------------

    def _detect_section_from_lines(self, lines: list[str]) -> Optional[str]:
        """
        Detecta a qué sección pertenece una fila según sus líneas de texto.
        Retorna "treasures", "ministry", "christian_life", o None.
        """
        for line in lines:
            upper = line.upper().strip().rstrip('.')
            if upper in ("TESOROS DE LA BIBLIA", "TESOROS DE LAS BIBLIAS"):
                return "treasures"
            if upper == "SEAMOS MEJORES MAESTROS":
                return "ministry"
            if upper == "NUESTRA VIDA CRISTIANA":
                return "christian_life"
        return None

    def _lines_contain_conclusion(self, lines: list[str]) -> bool:
        """Verifica si las líneas contienen 'Palabras de conclusión'."""
        for line in lines:
            if RE_CONCLUSION_WORDS.match(line):
                return True
        return False

    def _lines_contain_final_prayer(self, lines: list[str]) -> bool:
        """Verifica si las líneas contienen 'Oración final'."""
        for line in lines:
            if RE_FINAL_PRAYER.match(line):
                return True
        return False

    # ------------------------------------------------------------------
    # Creación de Assignment unificada
    # ------------------------------------------------------------------

    def _create_assignment(
        self, point_number: int, lines: list[str], person_text: str
    ) -> Assignment:
        """
        Crea un Assignment a partir de los datos de una fila.

        Estrategia corregida:
        1. Reconstruye el texto completo uniendo todas las líneas relevantes.
        2. Extrae la última duración del texto completo (regex tolerante).
        3. El resto del texto, incluido lo posterior a la duración, es el título.
        4. Extrae participantes desde person_text (columna separada).
        """
        import re as re_module

        assignment = Assignment(number=point_number)
        assignment.raw_lines = lines

        # ------------------------------------------------------------------
        # Paso 1: Detectar flags especiales (estudio bíblico, conclusión, oración)
        # ------------------------------------------------------------------
        is_bible_study = False
        is_conclusion = False
        is_prayer = False

        for line in lines:
            if RE_BIBLE_STUDY.match(line):
                is_bible_study = True
            if RE_CONCLUSION_WORDS.match(line):
                is_conclusion = True
            if RE_FINAL_PRAYER.match(line):
                is_prayer = True

        assignment.is_bible_study = is_bible_study
        assignment.is_conclusion = is_conclusion
        assignment.is_prayer = is_prayer

        # ------------------------------------------------------------------
        # Paso 2: Reconstruir texto completo de la asignación
        # Filtramos líneas que NO son contenido de la asignación:
        # encabezados de sección, canciones, oración final (ya procesada).
        # ------------------------------------------------------------------
        content_lines = []
        for line in lines:
            # Ignorar encabezados de sección
            if self._detect_section_from_lines([line]):
                continue
            # Ignorar canciones
            if RE_SONG.match(line):
                continue
            # Ignorar oración final (se procesa aparte)
            if RE_FINAL_PRAYER.match(line):
                continue
            # El resto es contenido
            cleaned = clean_text(line)
            if cleaned:
                content_lines.append(cleaned)

        # Unir todas las líneas de contenido con espacio
        full_text = " ".join(content_lines).strip()
        # Normalizar espacios múltiples
        full_text = re_module.sub(r'\s+', ' ', full_text)

        if not full_text:
            logger.debug("Assignment %d: sin contenido textual (solo encabezado).", point_number)
            if person_text:
                assignment.participants = self._parse_participants(person_text, is_bible_study)
            return assignment

        # ------------------------------------------------------------------
        # Paso 3: Extraer la última duración del texto completo.
        # ------------------------------------------------------------------
        # Buscar la ÚLTIMA ocurrencia de duración en CUALQUIER parte del texto
        # (no solo al final, porque puede haber texto después como "(parte 1)")
        dur_pattern = re_module.compile(
            r'[\(\.]?\s*(\d+)\s*(?:minutos?|mins?\.?)\s*\)?\.?',
            re_module.IGNORECASE,
        )
        # Encontrar TODAS las ocurrencias y quedarnos con la última
        all_dur_matches = list(dur_pattern.finditer(full_text))
        dur_match = all_dur_matches[-1] if all_dur_matches else None

        if dur_match:
            duration_minutes = int(dur_match.group(1))
            assignment.duration_minutes = duration_minutes
            assignment.duration_text = f"{duration_minutes} mins."

            # Se conserva cualquier aclaración escrita después de la duración.
            title = clean_text(
                full_text[:dur_match.start()] + full_text[dur_match.end():]
            ).strip("., ")
            assignment.title = title
        else:
            # Sin duración encontrada: todo el texto es el título
            assignment.title = full_text

        # ------------------------------------------------------------------
        # Paso 4: Validación
        # ------------------------------------------------------------------
        if not assignment.title.strip():
            logger.warning(
                "Assignment %d (BS=%s, Concl=%s): título vacío. "
                "Full text: '%s', Lines: %s",
                point_number, is_bible_study, is_conclusion,
                full_text, lines,
            )
        if not assignment.duration_text:
            logger.debug(
                "Assignment %d: sin duración detectada en '%s'",
                point_number, full_text[:60],
            )

        # ------------------------------------------------------------------
        # Paso 5: Extraer participantes (columna separada, ya funciona bien)
        # ------------------------------------------------------------------
        if person_text:
            assignment.participants = self._parse_participants(
                person_text, is_bible_study
            )

        return assignment

    def _parse_participants(
        self, person_text: str, is_bible_study: bool = False
    ) -> list[Participant]:
        """
        Parsea los participantes desde el texto de la columna 2.

        Divide por '//' y asigna roles automáticamente.
        """
        participants: list[Participant] = []

        if not person_text:
            return participants

        # Usar split_names que ya maneja //
        name1, name2 = split_names(person_text)

        if is_bible_study:
            # Roles: Conductor / Lector
            if name1:
                participants.append(Participant(name=name1, role="Conductor"))
            if name2:
                participants.append(Participant(name=name2, role="Lector"))
        else:
            # Roles genéricos: Estudiante / Ayudante
            if name1:
                participants.append(Participant(name=name1, role="Estudiante"))
            if name2:
                participants.append(Participant(name=name2, role="Ayudante"))

        # Si no se detectó // pero hay nombre, es un solo participante
        if not participants and person_text.strip():
            cleaned = normalize_person_name(remove_name_suffix(person_text))
            if cleaned:
                participants.append(Participant(name=cleaned, role=""))

        return participants

    # ------------------------------------------------------------------
    # Utilidades de la tabla
    # ------------------------------------------------------------------

    def _clean_point_number(self, text: str) -> Optional[int]:
        """Limpia el texto de número de punto."""
        if not text:
            return None
        cleaned = text.replace("//", "").strip()
        return extract_number(cleaned)

    def _split_cell_lines(self, cell_text: str) -> list[str]:
        """Divide el texto de una celda en líneas significativas."""
        lines = cell_text.split("\n")
        result = []
        for line in lines:
            cleaned = clean_text(line)
            if cleaned:
                result.append(cleaned)
        return result

    def _process_special_row(
        self, week: MeetingWeek, lines: list[str], person_text: str
    ) -> None:
        """Procesa una fila sin número de punto (solo encabezados/canciones)."""
        for line in lines:
            song_match = RE_SONG.match(line)
            if song_match:
                song_num = song_match.group(1)
                if not week.intermediate_song:
                    week.intermediate_song = song_num
                elif not week.closing_song:
                    week.closing_song = song_num

    # ------------------------------------------------------------------
    # Extracción del cierre
    # ------------------------------------------------------------------

    def _extract_closing_from_table(self, week: MeetingWeek, table) -> None:
        """Extrae canción final y oración final del final de la tabla."""
        songs: list[str] = []
        for row in table.rows:
            cells = row.cells
            if len(cells) < 2:
                continue

            cell1_text = cells[1].text
            person_text = clean_text(cells[2].text) if len(cells) > 2 else ""

            lines = self._split_cell_lines(cell1_text)
            for line in lines:
                if RE_FINAL_PRAYER.match(line):
                    if person_text and not week.closing_prayer:
                        student, _ = split_names(person_text)
                        week.closing_prayer = normalize_person_name(
                            remove_name_suffix(student)
                        )
                    continue

                song_match = RE_SONG.match(line)
                if song_match:
                    songs.append(song_match.group(1))

        if songs:
            week.intermediate_song = songs[0]
        if len(songs) > 1:
            week.closing_song = songs[-1]


def parse_document(file_path: Path) -> list[MeetingWeek]:
    """Función de conveniencia para parsear un documento."""
    parser = ProgramParser(file_path)
    return parser.parse()
