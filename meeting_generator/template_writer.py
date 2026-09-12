"""
template_writer.py — Escritor profesional del documento S-140.

Arquitectura:
1. DESCUBRIMIENTO Y PREFLIGHT: localiza las copias y valida su estructura y
   capacidad antes de modificarlas.
2. ESCRITURA POR FILA: cada Assignment se escribe como una unidad completa
   en su fila correspondiente (número, título, duración, participantes).
3. VERIFICACIÓN: comprueba valores escritos y marcadores en cada copia usada.
4. LIMPIEZA FINAL: limpia solo las copias no utilizadas y guarda cuando todo
   el documento queda libre de marcadores.

Columnas: 0-6 = [FECHA], 7-13 = resto. Sala auxiliar (8-12) se deja vacía.
Auditorio principal = columna 13.
"""

import hashlib
import logging
import os
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional, TypedDict

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table

from .config import (
    CONGREGATION_NAME,
    DEFAULT_DURATIONS,
    MEETING_START_MINUTES,
    MINISTRY_TRANSITION_MINUTES,
    OPENING_PRAYER_DURATION_MINUTES,
    SONG_DURATION_MINUTES,
)
from .models import Assignment, MeetingWeek
from .results import GenerationResult
from .validation import GenerationValidationError

logger = logging.getLogger(__name__)

PLACEHOLDERS = (
    "[FECHA]",
    "[Nombre]",
    "[Título]",
    "[Nombre/Nombre]",
    "Canción [Número]",
    "[Número]",
    "[X mins.]",
    "[XX mins.]",
    "(X mins.)",
    "(XX mins.)",
    "[NOMBRE DE LA CONGREGACIÓN]",
)
TEMPLATE_MARKERS = (*PLACEHOLDERS, "LECTURA SEMANAL DE LA BIBLIA")

HEADER_TREASURES = "TESOROS DE LA BIBLIA"
HEADER_MINISTRY = "SEAMOS MEJORES MAESTROS"
HEADER_CHRISTIAN_LIFE = "NUESTRA VIDA CRISTIANA"
HEADER_CONCLUSION = "Palabras de conclusión"


@dataclass
class RowDescriptor:
    row_index: int
    number_cell: Optional[int] = None
    title_cell: Optional[int] = None
    duration_cell: Optional[int] = None
    participant_cell: Optional[int] = None
    writes_title: bool = False
    writes_duration: bool = False
    auxiliary_cells: list[int] = field(default_factory=list)

    def content_cells(self) -> list[int]:
        cells: list[int] = []
        for cell_index in (
            self.number_cell,
            self.title_cell,
            self.duration_cell,
            self.participant_cell,
        ):
            if cell_index is not None and cell_index not in cells:
                cells.append(cell_index)
        return cells


class TemplateStructure(TypedDict):
    fecha_row: Optional[int]
    presidente_name_row: Optional[int]
    consejero_name_row: Optional[int]
    cancion_inicio_row: Optional[int]
    oracion_inicio_name_row: Optional[int]
    introduction_row: Optional[int]
    bible_treasures_rows: list[RowDescriptor]
    ministry_rows: list[RowDescriptor]
    cv_normal_rows: list[RowDescriptor]
    bible_study_row: Optional[RowDescriptor]
    conclusion_row: Optional[int]
    cancion_intermedia_row: Optional[int]
    cancion_final_row: Optional[int]
    oracion_final_name_row: Optional[int]


@dataclass(frozen=True)
class TemplateCopy:
    table: Table
    table_index: int
    start_row: int
    end_row: int
    structure: TemplateStructure


class TemplateWriter:
    """Escritor profesional del documento S-140."""

    def __init__(self, template_path: Path):
        self.template_path = template_path
        self.document: Optional[DocxDocument] = None

    # ==================================================================
    # API pública
    # ==================================================================

    def fill(
        self,
        weeks: list[MeetingWeek],
        output_path: Path,
    ) -> GenerationResult:
        """Ejecuta la generación completa y devuelve un resultado verificable."""
        try:
            if self.template_path.resolve() == output_path.resolve():
                raise GenerationValidationError(
                    "Ruta de salida inválida",
                    ["La ruta de salida coincide con la plantilla"],
                )
            created_path = self._fill_document(weeks, output_path)
        except GenerationValidationError as exc:
            logger.error("Generación rechazada: %s", exc)
            return GenerationResult.failure(len(weeks), list(exc.issues))
        except Exception as exc:
            logger.error("La generación falló", exc_info=True)
            return GenerationResult.failure(
                len(weeks),
                [f"No se pudo generar el documento: {exc}"],
            )

        return GenerationResult.success(len(weeks), created_path)

    def _fill_document(self, weeks: list[MeetingWeek], output_path: Path) -> Path:
        logger.info("Cargando plantilla: %s", self.template_path)
        self.document = Document(str(self.template_path))

        copies = self._discover_template_copies()
        logger.info("Capacidad detectada en la plantilla: %d formularios.", len(copies))
        issues = self._input_validation_issues(weeks, copies)
        if issues:
            raise GenerationValidationError("No se puede generar el documento", issues)

        write_issues: list[str] = []
        for week, copy in zip(weeks, copies):
            self._process_one_copy(
                copy.table,
                copy.start_row,
                copy.end_row,
                week,
                copy.structure,
            )
            write_issues.extend(self._written_content_issues(week, copy))

        remaining = self._clear_unused_copies_and_find_markers(copies, len(weeks))
        if remaining:
            write_issues.append(f"sin resolver: {', '.join(remaining)}")
        if write_issues:
            raise GenerationValidationError("La plantilla no se pudo completar", write_issues)
        logger.info("Marcadores residuales tras la validación: 0")

        created_path = self._save_document_atomically(output_path)
        logger.info("Documento guardado: %s", created_path)
        return created_path

    def _save_document_atomically(self, output_path: Path) -> Path:
        """Publica un DOCX verificado sin dejar una salida parcial."""
        if self.document is None:
            raise GenerationValidationError(
                "No se puede guardar el documento",
                ["no hay un documento preparado para guardar"],
            )
        if not output_path.parent.is_dir():
            raise GenerationValidationError(
                "No se puede guardar el documento",
                [f"el directorio de salida no existe: {output_path.parent}"],
            )

        descriptor, temporary_name = tempfile.mkstemp(
            dir=str(output_path.parent),
            prefix=f".{output_path.stem}-",
            suffix=".docx",
        )
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        try:
            self.document.save(str(temporary_path))
            if not temporary_path.is_file() or temporary_path.stat().st_size <= 0:
                raise OSError("el archivo temporal no se creó correctamente")
            Document(str(temporary_path))
            expected_digest = self._file_digest(temporary_path)
            os.replace(temporary_path, output_path)
            created_path = output_path.resolve(strict=True)
            if not created_path.is_file() or created_path.stat().st_size <= 0:
                raise OSError("no se pudo verificar el archivo de salida")
            if self._file_digest(created_path) != expected_digest:
                raise OSError("el archivo publicado no coincide con el generado")
            return created_path
        finally:
            temporary_path.unlink(missing_ok=True)

    def _file_digest(self, path: Path) -> bytes:
        digest = hashlib.sha256()
        with path.open("rb") as document_file:
            for chunk in iter(lambda: document_file.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.digest()

    # ==================================================================
    # Descubrimiento y validación de formularios
    # ==================================================================

    def _discover_template_copies(self) -> list[TemplateCopy]:
        if self.document is None:
            return []

        copies: list[TemplateCopy] = []
        for table_index, table in enumerate(self.document.tables, start=1):
            starts = self._find_copy_starts(table)
            for copy_index, start_row in enumerate(starts):
                end_row = (
                    starts[copy_index + 1]
                    if copy_index + 1 < len(starts)
                    else len(table.rows)
                )
                copies.append(
                    TemplateCopy(
                        table=table,
                        table_index=table_index,
                        start_row=start_row,
                        end_row=end_row,
                        structure=self._discover_structure(table, start_row, end_row),
                    )
                )
        return copies

    def _input_validation_issues(
        self,
        weeks: list[MeetingWeek],
        copies: list[TemplateCopy],
    ) -> list[str]:
        issues: list[str] = []
        if not weeks:
            issues.append("no hay semanas para escribir")
        if not copies:
            issues.append("la plantilla no contiene formularios reconocibles")
        if len(weeks) > len(copies):
            issues.append(
                f"la capacidad de la plantilla es de {len(copies)} semanas, "
                f"pero se recibieron {len(weeks)}"
            )

        for week in weeks:
            if not week.is_valid():
                issues.extend(
                    f"semana {week.week_index}: {issue}"
                    for issue in week.validation_errors()
                )

        for week, copy in zip(weeks, copies):
            issues.extend(self._copy_validation_issues(week, copy))
        return issues

    def _copy_validation_issues(
        self,
        week: MeetingWeek,
        copy: TemplateCopy,
    ) -> list[str]:
        prefix = f"semana {week.week_index}, tabla {copy.table_index}"
        issues: list[str] = []
        structure = copy.structure

        required_rows = (
            ("fecha", structure["fecha_row"]),
            ("presidente", structure["presidente_name_row"]),
            ("canción inicial", structure["cancion_inicio_row"]),
            ("oración inicial", structure["oracion_inicio_name_row"]),
            ("introducción", structure["introduction_row"]),
            ("canción intermedia", structure["cancion_intermedia_row"]),
            ("estudio bíblico", structure["bible_study_row"]),
            ("conclusión", structure["conclusion_row"]),
            ("canción final", structure["cancion_final_row"]),
            ("oración final", structure["oracion_final_name_row"]),
        )
        for field_name, row in required_rows:
            if row is None:
                issues.append(f"{prefix}: falta la fila de {field_name}")

        required_labels = (
            "LECTURA SEMANAL DE LA BIBLIA",
            HEADER_TREASURES,
            HEADER_MINISTRY,
            HEADER_CHRISTIAN_LIFE,
            HEADER_CONCLUSION,
        )
        for label in required_labels:
            if not self._range_contains(copy, label):
                issues.append(f"{prefix}: falta el marcador '{label}'")

        christian_life_normal = [
            assignment
            for assignment in week.christian_life
            if not assignment.is_bible_study and not assignment.is_conclusion
        ]
        capacities = (
            (
                "Tesoros de la Biblia",
                week.bible_treasures,
                structure["bible_treasures_rows"],
            ),
            ("Seamos Mejores Maestros", week.ministry, structure["ministry_rows"]),
            (
                "Nuestra Vida Cristiana",
                christian_life_normal,
                structure["cv_normal_rows"],
            ),
        )
        for section_name, assignments, descriptors in capacities:
            if len(assignments) > len(descriptors):
                issues.append(
                    f"{prefix}: la capacidad de {section_name} es de "
                    f"{len(descriptors)} filas y se necesitan {len(assignments)}"
                )
            for assignment, descriptor in zip(assignments, descriptors):
                issues.extend(
                    self._descriptor_validation_issues(
                        prefix,
                        section_name,
                        copy.table,
                        descriptor,
                        assignment,
                    )
                )

        bible_study_row = structure["bible_study_row"]
        if bible_study_row is not None:
            bible_study = next(
                (
                    assignment
                    for assignment in week.christian_life
                    if assignment.is_bible_study
                ),
                None,
            )
        else:
            bible_study = None
        if bible_study_row is not None and bible_study is not None:
            issues.extend(
                self._descriptor_validation_issues(
                    prefix,
                    "Estudio bíblico",
                    copy.table,
                    bible_study_row,
                    bible_study,
                )
            )

        conclusion_row = structure["conclusion_row"]
        conclusions = [
            assignment
            for assignment in week.christian_life
            if assignment.is_conclusion
        ]
        if conclusion_row is not None:
            template_duration = self._fixed_duration_in_row(
                copy.table,
                conclusion_row,
            )
            expected_duration = DEFAULT_DURATIONS["conclusion"]
            if template_duration != expected_duration:
                issues.append(
                    f"{prefix}: la conclusión de la plantilla debe durar "
                    f"{expected_duration} minutos"
                )
            if (
                len(conclusions) == 1
                and template_duration is not None
                and conclusions[0].duration_minutes != template_duration
            ):
                issues.append(
                    f"{prefix}: la duración de la conclusión no coincide con la "
                    "plantilla"
                )
        return issues

    def _descriptor_validation_issues(
        self,
        prefix: str,
        section_name: str,
        table: Table,
        descriptor: RowDescriptor,
        assignment: Assignment,
    ) -> list[str]:
        row = table.rows[descriptor.row_index]
        texts = self._unique_cell_texts(row)
        missing: list[str] = []
        if not any(self._starts_with_number(text) for text in texts):
            missing.append("número")
        if not any(self._has_writable_title(text) for text in texts):
            missing.append("título")
        if not any(
            self._has_duration_marker(text)
            or self._fixed_duration_from_text(text) is not None
            for text in texts
        ):
            missing.append("duración")
        if descriptor.participant_cell is None or not any(
            "[Nombre]" in text or "[Nombre/Nombre]" in text for text in texts
        ):
            missing.append("participantes")

        issues = []
        if missing:
            issues.append(
                f"{prefix}: una fila de {section_name} no admite {', '.join(missing)}"
            )

        fixed_duration = next(
            (
                duration
                for text in texts
                if (duration := self._fixed_duration_from_text(text)) is not None
            ),
            None,
        )
        if fixed_duration is not None and assignment.duration_minutes != fixed_duration:
            issues.append(
                f"{prefix}: el punto {assignment.number} dura "
                f"{assignment.duration_minutes} minutos, pero su fila fija "
                f"muestra {fixed_duration}"
            )
        return issues

    def _unique_cell_texts(self, row) -> list[str]:
        texts: list[str] = []
        seen_cells: set[object] = set()
        for cell in row.cells:
            if cell._tc in seen_cells:
                continue
            seen_cells.add(cell._tc)
            texts.append(cell.text.strip())
        return texts

    def _has_writable_title(self, text: str) -> bool:
        if "[Título]" in text:
            return True
        prefix = self._leading_number_prefix(text)
        if not prefix:
            return False
        duration_match = re.search(r"\(\s*\d+\s+mins?\.?\s*\)", text, re.IGNORECASE)
        if duration_match is None:
            return False
        return bool(text[len(prefix) : duration_match.start()].strip())

    def _fixed_duration_from_text(self, text: str) -> Optional[int]:
        match = re.search(r"\(\s*(\d+)\s+mins?\.?\s*\)", text, re.IGNORECASE)
        return int(match.group(1)) if match else None

    def _fixed_duration_in_row(self, table: Table, row_index: int) -> Optional[int]:
        for text in self._unique_cell_texts(table.rows[row_index]):
            duration = self._fixed_duration_from_text(text)
            if duration is not None:
                return duration
        return None

    def _written_content_issues(
        self,
        week: MeetingWeek,
        copy: TemplateCopy,
    ) -> list[str]:
        """Comprueba que los datos esperados estén presentes tras escribir."""
        issues: list[str] = []
        structure = copy.structure
        prefix = f"semana {week.week_index}, tabla {copy.table_index}"

        row_values = (
            ("fecha", structure["fecha_row"], week.date),
            ("lectura semanal", structure["fecha_row"], week.weekly_reading),
            ("presidente", structure["presidente_name_row"], week.president),
            (
                "oración inicial",
                structure["oracion_inicio_name_row"],
                week.opening_prayer,
            ),
            (
                "oración final",
                structure["oracion_final_name_row"],
                week.closing_prayer,
            ),
        )
        for field_name, row_index, expected_value in row_values:
            if row_index is not None and not self._row_contains(
                copy.table,
                row_index,
                expected_value,
            ):
                issues.append(f"{prefix}: no se escribió {field_name}")

        songs = (
            ("canción inicial", structure["cancion_inicio_row"], week.opening_song),
            (
                "canción intermedia",
                structure["cancion_intermedia_row"],
                week.intermediate_song,
            ),
            ("canción final", structure["cancion_final_row"], week.closing_song),
        )
        for field_name, row_index, song_number in songs:
            if row_index is not None and not self._row_contains_pattern(
                copy.table,
                row_index,
                rf"Canción\s+{re.escape(song_number)}(?!\d)",
            ):
                issues.append(f"{prefix}: no se escribió {field_name}")

        christian_life_normal = [
            assignment
            for assignment in week.christian_life
            if not assignment.is_bible_study and not assignment.is_conclusion
        ]
        sections = (
            (week.bible_treasures, structure["bible_treasures_rows"]),
            (week.ministry, structure["ministry_rows"]),
            (christian_life_normal, structure["cv_normal_rows"]),
        )
        for assignments, descriptors in sections:
            for assignment, descriptor in zip(
                sorted(assignments, key=lambda item: item.number),
                descriptors,
            ):
                issues.extend(
                    self._assignment_written_issues(
                        prefix,
                        copy.table,
                        descriptor,
                        assignment,
                    )
                )

        bible_study = next(
            (
                assignment
                for assignment in week.christian_life
                if assignment.is_bible_study
            ),
            None,
        )
        bible_study_row = structure["bible_study_row"]
        if bible_study is not None and bible_study_row is not None:
            issues.extend(
                self._assignment_written_issues(
                    prefix,
                    copy.table,
                    bible_study_row,
                    bible_study,
                )
            )
        return issues

    def _assignment_written_issues(
        self,
        prefix: str,
        table: Table,
        descriptor: RowDescriptor,
        assignment: Assignment,
    ) -> list[str]:
        issues: list[str] = []
        row = table.rows[descriptor.row_index]
        label = f"punto {assignment.number}"

        number_cell = descriptor.number_cell
        if number_cell is not None and not self._cell_starts_with_number(
            row.cells[number_cell],
            assignment.number,
        ):
            issues.append(f"{prefix}: no se escribió el número del {label}")

        title_cell = descriptor.title_cell
        if (
            descriptor.writes_title
            and title_cell is not None
            and assignment.title not in row.cells[title_cell].text
        ):
            issues.append(f"{prefix}: no se escribió el título del {label}")

        duration_cell = descriptor.duration_cell
        if (
            descriptor.writes_duration
            and duration_cell is not None
            and assignment.duration_text not in row.cells[duration_cell].text
        ):
            issues.append(f"{prefix}: no se escribió la duración del {label}")

        participant_cell = descriptor.participant_cell
        participants = assignment.formatted_participants()
        if (
            participant_cell is not None
            and participants not in row.cells[participant_cell].text
        ):
            issues.append(f"{prefix}: no se escribieron los participantes del {label}")
        return issues

    def _cell_starts_with_number(self, cell, expected_number: int) -> bool:
        expected_prefix = f"{expected_number}."
        return any(
            self._leading_number_prefix(paragraph.text) == expected_prefix
            for paragraph in cell.paragraphs
        )

    def _row_contains(self, table: Table, row_index: int, text: str) -> bool:
        return any(
            text in cell_text
            for cell_text in self._unique_cell_texts(table.rows[row_index])
        )

    def _row_contains_pattern(
        self,
        table: Table,
        row_index: int,
        pattern: str,
    ) -> bool:
        return any(
            re.search(pattern, cell_text) is not None
            for cell_text in self._unique_cell_texts(table.rows[row_index])
        )

    def _range_contains(self, copy: TemplateCopy, text: str) -> bool:
        for row_index in range(copy.start_row, copy.end_row):
            if any(text in cell.text for cell in copy.table.rows[row_index].cells):
                return True
        return False

    def _find_copy_starts(self, table) -> list[int]:
        starts = []
        for ri, row in enumerate(table.rows):
            if len(row.cells) > 0 and "[FECHA]" in row.cells[0].text:
                starts.append(ri)
        return starts

    # ==================================================================
    # Procesamiento de UNA copia del formulario
    # ==================================================================

    def _process_one_copy(
        self,
        table,
        start_row: int,
        end_row: int,
        week: MeetingWeek,
        structure: TemplateStructure,
    ) -> None:
        """Rellena una copia del formulario."""
        # 1. Cabecera fija
        self._write_header(table, week, structure)

        # 2. Tesoros de la Biblia
        self._write_section_rows(table, week.bible_treasures, structure["bible_treasures_rows"])

        # 3. Ministerio — escribir los existentes, limpiar sobrantes
        self._write_section_rows(table, week.ministry, structure["ministry_rows"])

        # 4. Vida Cristiana (solo normales, no BS ni conclusión)
        cv_normal = [a for a in week.christian_life
                     if not a.is_bible_study and not a.is_conclusion]
        self._write_section_rows(table, cv_normal, structure["cv_normal_rows"])

        # 5. Estudio bíblico
        bible_study_row = structure["bible_study_row"]
        if bible_study_row is not None:
            eb = next((a for a in week.christian_life if a.is_bible_study), None)
            if eb:
                self._write_bible_study_row(table, bible_study_row, eb)
            else:
                self._clear_row(
                    table,
                    bible_study_row.row_index,
                )

        # 6. Canciones y oraciones
        self._write_songs_and_prayers(table, week, structure)

        # 7. Horas de inicio
        self._write_schedule_times(table, start_row, end_row, week, structure)

    # ==================================================================
    # Descubrimiento de estructura
    # ==================================================================

    def _discover_structure(
        self, table, start_row: int, end_row: int
    ) -> TemplateStructure:
        """Escanea las filas y descubre la posición de cada elemento."""
        struct: TemplateStructure = {
            "fecha_row": None,
            "presidente_name_row": None,
            "consejero_name_row": None,
            "cancion_inicio_row": None,
            "oracion_inicio_name_row": None,
            "introduction_row": None,
            "bible_treasures_rows": [],
            "ministry_rows": [],
            "cv_normal_rows": [],
            "bible_study_row": None,
            "conclusion_row": None,
            "cancion_intermedia_row": None,
            "cancion_final_row": None,
            "oracion_final_name_row": None,
        }

        current_section = "header"

        for ri in range(start_row, end_row):
            row = table.rows[ri]
            combined = " | ".join(
                row.cells[ci].text.strip()
                for ci in range(min(14, len(row.cells)))
            )

            if HEADER_MINISTRY in combined.upper() and current_section != "christian_life":
                current_section = "ministry"
                continue
            if HEADER_CHRISTIAN_LIFE in combined.upper():
                current_section = "christian_life"
                continue

            if current_section == "header":
                if "[FECHA]" in combined:
                    struct["fecha_row"] = ri
                if "Presidente:" in combined:
                    struct["presidente_name_row"] = self._find_nombre_row(table, ri, end_row)
                if "Consejero" in combined and struct["consejero_name_row"] is None:
                    for ci in range(7, 14):
                        if ci < len(row.cells) and "[Nombre]" in row.cells[ci].text:
                            struct["consejero_name_row"] = ri
                            break
                if "Canción [Número]" in combined and struct["cancion_inicio_row"] is None:
                    struct["cancion_inicio_row"] = ri
                if "Oración:" in combined and struct["oracion_inicio_name_row"] is None:
                    struct["oracion_inicio_name_row"] = self._find_nombre_in_row(row, ri)
                if "Palabras de introducción" in combined:
                    struct["introduction_row"] = ri
                if (
                    ("[Título]" in combined or "[Nombre/Nombre]" in combined or "[Nombre]" in combined)
                    and "Presidente:" not in combined
                    and "Consejero" not in combined
                    and "Oración:" not in combined
                    and "Canción [Número]" not in combined
                ):
                    struct["bible_treasures_rows"].append(self._build_row_descriptor(table, ri))

            elif current_section == "ministry":
                if "[Título]" in combined or "[Nombre/Nombre]" in combined:
                    struct["ministry_rows"].append(self._build_row_descriptor(table, ri))

            elif current_section == "christian_life":
                if "Canción [Número]" in combined and struct["cancion_intermedia_row"] is None:
                    struct["cancion_intermedia_row"] = ri
                    continue
                if "Estudio bíblico" in combined:
                    struct["bible_study_row"] = self._build_row_descriptor(table, ri)
                    continue
                if "[Título]" in combined or "[Nombre]" in combined:
                    struct["cv_normal_rows"].append(self._build_row_descriptor(table, ri))
                    continue
                if HEADER_CONCLUSION in combined:
                    struct["conclusion_row"] = ri
                    current_section = "closing"
                    continue

            elif current_section == "closing":
                if "Canción [Número]" in combined and struct["cancion_final_row"] is None:
                    struct["cancion_final_row"] = ri
                if "Oración:" in combined and struct["oracion_final_name_row"] is None:
                    struct["oracion_final_name_row"] = self._find_nombre_in_row(row, ri)

        return struct

    def _build_row_descriptor(self, table, row_idx: int) -> RowDescriptor:
        row = table.rows[row_idx]
        descriptor = RowDescriptor(row_index=row_idx)
        participant_cells_by_identity: dict[object, int] = {}

        for ci, cell in enumerate(row.cells):
            text = cell.text.strip()
            if not text:
                continue

            if text == "Sala auxiliar":
                descriptor.auxiliary_cells.append(ci)

            if "[Nombre/Nombre]" in text or "[Nombre]" in text:
                # Una celda física puede repetirse en row.cells por gridSpan; conservamos
                # la última posición vista de cada tc para escoger luego la de mayor índice.
                participant_cells_by_identity[cell._tc] = ci

            if descriptor.title_cell is None and (
                "[Título]" in text or "Estudio bíblico" in text
            ):
                descriptor.title_cell = ci
                descriptor.duration_cell = ci
                descriptor.writes_title = "[Título]" in text
                descriptor.writes_duration = self._has_duration_marker(text)
                if descriptor.number_cell is None and self._starts_with_number(text):
                    descriptor.number_cell = ci
                continue

            if descriptor.number_cell is None and self._starts_with_number(text):
                descriptor.number_cell = ci

            if descriptor.duration_cell is None and self._has_duration_marker(text):
                descriptor.duration_cell = ci

            if "[Título]" in text:
                descriptor.writes_title = True
            if self._has_duration_marker(text):
                descriptor.writes_duration = True

        if descriptor.title_cell is None:
            descriptor.title_cell = descriptor.number_cell
        if descriptor.duration_cell is None:
            descriptor.duration_cell = descriptor.title_cell
        if descriptor.number_cell is None:
            descriptor.number_cell = descriptor.title_cell

        if participant_cells_by_identity:
            participant_indices = sorted(participant_cells_by_identity.values())
            descriptor.participant_cell = participant_indices[-1]
            descriptor.auxiliary_cells.extend(
                ci for ci in participant_indices[:-1] if ci not in descriptor.auxiliary_cells
            )

        logger.debug(
            "Descriptor row=%s number_cell=%s title_cell=%s duration_cell=%s participant_cell=%s auxiliary_cells=%s",
            descriptor.row_index,
            descriptor.number_cell,
            descriptor.title_cell,
            descriptor.duration_cell,
            descriptor.participant_cell,
            descriptor.auxiliary_cells,
        )
        return descriptor

    def _starts_with_number(self, text: str) -> bool:
        stripped = text.lstrip()
        digits: list[str] = []
        for char in stripped:
            if char.isdigit():
                digits.append(char)
                continue
            break
        if not digits:
            return False
        remaining = stripped[len(digits):].lstrip()
        return remaining.startswith(".")

    def _has_duration_marker(self, text: str) -> bool:
        return any(marker in text for marker in ("[X mins.]", "[XX mins.]", "(X mins.)", "(XX mins.)"))

    def _find_nombre_row(self, table, ri: int, end_row: int) -> Optional[int]:
        """Busca la fila con [Nombre] a partir de ri (puede ser la misma o siguiente)."""
        for offset in range(2):
            check = ri + offset
            if check < end_row and check < len(table.rows):
                for ci in range(7, 14):
                    if ci < len(table.rows[check].cells):
                        if "[Nombre]" in table.rows[check].cells[ci].text:
                            return check
        return None

    def _find_nombre_in_row(self, row, row_idx: int) -> Optional[int]:
        """Busca [Nombre] en la misma fila, retorna el índice de fila o None."""
        for ci in range(7, 14):
            if ci < len(row.cells) and "[Nombre]" in row.cells[ci].text:
                return row_idx
        return None

    # ==================================================================
    # Escritura de cabecera fija
    # ==================================================================

    def _write_header(
        self, table, week: MeetingWeek, struct: TemplateStructure
    ) -> None:
        """Escribe los campos fijos de la cabecera."""
        # [FECHA]
        if struct["fecha_row"] is not None and week.date:
            self._set_cell_text(table, struct["fecha_row"], 0, week.date, "[FECHA]")

        # Lectura semanal
        if struct["fecha_row"] is not None and week.weekly_reading:
            self._set_cell_text(
                table,
                struct["fecha_row"],
                0,
                week.weekly_reading,
                "LECTURA SEMANAL DE LA BIBLIA",
            )

        # Presidente
        if struct["presidente_name_row"] is not None and week.president:
            self._write_nombre_in_row(table, struct["presidente_name_row"], week.president)

        # Consejero
        if struct["consejero_name_row"] is not None:
            self._write_nombre_in_row(table, struct["consejero_name_row"], week.president)

        # Canción inicial
        if struct["cancion_inicio_row"] is not None and week.opening_song:
            self._write_cancion_in_row(table, struct["cancion_inicio_row"], week.opening_song)

        # Oración inicial
        if struct["oracion_inicio_name_row"] is not None:
            name = week.opening_prayer or week.president
            self._write_nombre_in_row(table, struct["oracion_inicio_name_row"], name)

    # ==================================================================
    # Escritura de secciones dinámicas
    # ==================================================================

    def _write_section_rows(self, table, assignments: list[Assignment],
                              row_descriptors: list[RowDescriptor]) -> None:
        """Escribe asignaciones en las filas existentes; limpia las sobrantes."""
        ordered_assignments = sorted(assignments, key=lambda assignment: assignment.number)
        for i, descriptor in enumerate(row_descriptors):
            if descriptor.row_index >= len(table.rows):
                continue
            if i < len(ordered_assignments):
                self._write_assignment_in_row(table, descriptor, ordered_assignments[i])
            else:
                self._clear_row(table, descriptor.row_index)
                self._clear_time_in_row(table, descriptor.row_index)

    def _write_assignment_in_row(self, table, descriptor: RowDescriptor, assignment: Assignment) -> None:
        """Escribe una asignación completa en una fila."""
        row = table.rows[descriptor.row_index]

        self._clear_auxiliary_cells(row, descriptor)

        participant_cell = descriptor.participant_cell

        for ci, cell in enumerate(row.cells):
            text = cell.text
            if "[Número]" in text or self._starts_with_number(text):
                self._write_number_in_cell(cell, assignment.number)

            if "[Título]" in text:
                self._write_title_in_cell(cell, assignment.title)

            if self._has_duration_marker(text):
                self._write_duration_in_cell(cell, assignment.duration_text)

            if "[Nombre/Nombre]" in text or "[Nombre]" in text:
                if participant_cell is not None and ci == participant_cell:
                    # La celda principal es la de mayor índice entre las que contienen
                    # el placeholder de participante; las anteriores se vacían.
                    logger.debug(
                        "Fila %s: participantes escritos en columna %s",
                        descriptor.row_index,
                        ci,
                    )
                    self._write_participants_in_cell(cell, assignment)
                else:
                    logger.debug(
                        "Fila %s: celda auxiliar vaciada en columna %s",
                        descriptor.row_index,
                        ci,
                    )
                    self._empty_cell(cell)

    def _write_bible_study_row(self, table, descriptor: RowDescriptor, assignment: Assignment) -> None:
        """Escribe el estudio bíblico."""
        self._write_assignment_in_row(table, descriptor, assignment)

    def _clear_auxiliary_cells(self, row, descriptor: RowDescriptor) -> None:
        seen_cells: set[object] = set()
        for ci in descriptor.auxiliary_cells:
            if ci < len(row.cells):
                cell = row.cells[ci]
                if cell._tc in seen_cells:
                    continue
                seen_cells.add(cell._tc)
                self._empty_cell(cell)

    def _empty_cell(self, cell) -> None:
        for par in cell.paragraphs:
            for run in par.runs:
                run.text = ""

    def _write_number_in_cell(self, cell, number: int) -> None:
        if self._replace_in_cell_conditional(cell, "[Número]", str(number)):
            return
        for par in cell.paragraphs:
            full = self._paragraph_text(par)
            prefix = self._leading_number_prefix(full)
            if prefix:
                self._replace_in_paragraph(par, prefix, f"{number}.")
                return

    def _write_title_in_cell(self, cell, title: str) -> None:
        if not title:
            return
        if self._replace_in_cell_conditional(cell, "[Título]", title):
            return
        for par in cell.paragraphs:
            full = self._paragraph_text(par)
            if title in full:
                return
            prefix = self._leading_number_prefix(full)
            if not prefix:
                continue
            duration_index = self._first_duration_marker_index(full)
            if duration_index == -1:
                continue
            middle = full[len(prefix):duration_index]
            search = f"{prefix}{middle}"
            replacement = f"{prefix} {title} "
            if self._replace_in_paragraph(par, search, replacement):
                return

    def _write_duration_in_cell(self, cell, duration_text: str) -> None:
        if not duration_text:
            return
        for placeholder in ("[X mins.]", "[XX mins.]", "(X mins.)", "(XX mins.)"):
            replacement = duration_text if placeholder.startswith("[") else f"({duration_text})"
            if self._replace_in_cell_conditional(cell, placeholder, replacement):
                return

    def _write_participants_in_cell(self, cell, assignment: Assignment) -> None:
        participants = assignment.formatted_participants()
        self._replace_in_cell_conditional(cell, "[Nombre/Nombre]", participants)
        self._replace_in_cell_conditional(cell, "[Nombre]", participants)

    # ==================================================================
    # Canciones y oraciones
    # ==================================================================

    def _write_songs_and_prayers(
        self, table, week: MeetingWeek, struct: TemplateStructure
    ) -> None:
        # Canción intermedia
        if struct["cancion_intermedia_row"] is not None and week.intermediate_song:
            self._write_cancion_in_row(table, struct["cancion_intermedia_row"],
                                        week.intermediate_song)
        # Canción final
        if struct["cancion_final_row"] is not None and week.closing_song:
            self._write_cancion_in_row(table, struct["cancion_final_row"], week.closing_song)
        # Oración final
        if struct["oracion_final_name_row"] is not None and week.closing_prayer:
            self._write_nombre_in_row(table, struct["oracion_final_name_row"],
                                       week.closing_prayer)

    # ==================================================================
    # Helpers de escritura en fila
    # ==================================================================

    def _write_nombre_in_row(self, table, row_idx: int, name: str) -> None:
        if row_idx >= len(table.rows):
            return
        for ci in range(7, 14):
            if ci < len(table.rows[row_idx].cells):
                cell = table.rows[row_idx].cells[ci]
                if "[Nombre]" in cell.text and "[Nombre/Nombre]" not in cell.text:
                    self._replace_in_cell(cell, "[Nombre]", name)
                    return

    def _write_cancion_in_row(self, table, row_idx: int, song: str) -> None:
        if row_idx >= len(table.rows):
            return
        seen_cells: set[object] = set()
        for cell in table.rows[row_idx].cells:
            if cell._tc in seen_cells:
                continue
            seen_cells.add(cell._tc)
            if "Canción [Número]" in cell.text:
                self._replace_in_cell(cell, "Canción [Número]", f"Canción {song}")
                return

    # ==================================================================
    # Cálculo de horas
    # ==================================================================

    def _write_schedule_times(
        self,
        table,
        start_row: int,
        end_row: int,
        week: MeetingWeek,
        struct: TemplateStructure,
    ) -> None:
        """Calcula y escribe la hora de inicio de cada actividad."""
        events: list[tuple[int, int]] = []

        self._add_schedule_event(
            events,
            struct.get("cancion_inicio_row"),
            SONG_DURATION_MINUTES + OPENING_PRAYER_DURATION_MINUTES,
        )
        self._add_schedule_event(
            events,
            struct.get("introduction_row"),
            DEFAULT_DURATIONS["introduccion"],
        )
        self._add_assignment_events(
            events,
            week.bible_treasures,
            struct["bible_treasures_rows"],
        )
        self._add_assignment_events(
            events,
            week.ministry,
            struct["ministry_rows"],
            additional_minutes=MINISTRY_TRANSITION_MINUTES,
        )
        self._add_schedule_event(
            events,
            struct.get("cancion_intermedia_row"),
            SONG_DURATION_MINUTES,
        )

        christian_life = [
            assignment
            for assignment in week.christian_life
            if not assignment.is_bible_study and not assignment.is_conclusion
        ]
        self._add_assignment_events(
            events,
            christian_life,
            struct["cv_normal_rows"],
        )

        bible_study = next(
            (
                assignment
                for assignment in week.christian_life
                if assignment.is_bible_study
            ),
            None,
        )
        if bible_study is not None:
            self._add_schedule_event(
                events,
                self._descriptor_row(struct.get("bible_study_row")),
                self._assignment_duration(bible_study),
            )

        self._add_schedule_event(
            events,
            struct.get("conclusion_row"),
            DEFAULT_DURATIONS["conclusion"],
        )
        self._add_schedule_event(
            events,
            struct.get("cancion_final_row"),
            SONG_DURATION_MINUTES,
        )

        events.sort(key=lambda event: event[0])
        active_rows = {row_index for row_index, _ in events}
        for row_index in range(start_row, min(end_row, len(table.rows))):
            if row_index not in active_rows:
                self._clear_time_in_row(table, row_index)

        current_minutes = MEETING_START_MINUTES
        for row_index, duration_minutes in events:
            self._write_time_in_row(table, row_index, current_minutes)
            current_minutes += duration_minutes

    def _add_assignment_events(
        self,
        events: list[tuple[int, int]],
        assignments: list[Assignment],
        descriptors: list[RowDescriptor],
        additional_minutes: int = 0,
    ) -> None:
        ordered_assignments = sorted(assignments, key=lambda assignment: assignment.number)
        for assignment, descriptor in zip(ordered_assignments, descriptors):
            duration = self._assignment_duration(assignment) + additional_minutes
            self._add_schedule_event(events, descriptor.row_index, duration)

    def _assignment_duration(self, assignment: Assignment) -> int:
        if (
            isinstance(assignment.duration_minutes, int)
            and not isinstance(assignment.duration_minutes, bool)
            and assignment.duration_minutes > 0
        ):
            return assignment.duration_minutes
        raise GenerationValidationError(
            "No se puede calcular el horario",
            [f"el punto {assignment.number} no tiene una duración positiva"],
        )

    def _add_schedule_event(
        self,
        events: list[tuple[int, int]],
        row_index: Optional[int],
        duration_minutes: int,
    ) -> None:
        if row_index is not None:
            events.append((row_index, duration_minutes))

    def _descriptor_row(self, descriptor: Optional[RowDescriptor]) -> Optional[int]:
        return descriptor.row_index if descriptor is not None else None

    def _write_time_in_row(self, table, row_idx: int, total_minutes: int) -> None:
        if row_idx >= len(table.rows) or not table.rows[row_idx].cells:
            return
        cell = table.rows[row_idx].cells[0]
        current_text = cell.text.strip()
        display_text = self._format_time(total_minutes)
        if current_text:
            self._replace_in_cell(cell, current_text, display_text)

    def _clear_time_in_row(self, table, row_idx: int) -> None:
        if row_idx >= len(table.rows) or not table.rows[row_idx].cells:
            return
        cell = table.rows[row_idx].cells[0]
        if cell.text.strip() == "0:00":
            self._remove_text(cell, "0:00")

    def _format_time(self, total_minutes: int) -> str:
        hours, minutes = divmod(total_minutes, 60)
        display_hour = hours % 12 or 12
        return f"{display_hour}:{minutes:02d}"

    # ==================================================================
    # Limpieza de filas y celdas
    # ==================================================================

    def _clear_row(self, table, row_idx: int) -> None:
        if row_idx >= len(table.rows):
            return
        seen_cells: set[object] = set()
        for cell in table.rows[row_idx].cells:
            if cell._tc in seen_cells:
                continue
            seen_cells.add(cell._tc)
            text = cell.text
            for marker in PLACEHOLDERS:
                if marker in text:
                    self._remove_text(cell, marker)
                    text = cell.text

    def _clear_unused_copies_and_find_markers(
        self,
        copies: list[TemplateCopy],
        used_copy_count: int,
    ) -> list[str]:
        """Limpia y valida cada celda física de la plantilla en un solo pase."""
        if self.document is None:
            return []

        unused_ranges: dict[object, list[tuple[int, int]]] = {}
        for copy in copies[used_copy_count:]:
            unused_ranges.setdefault(copy.table._tbl, []).append(
                (copy.start_row, copy.end_row)
            )

        congregation_marker = "[NOMBRE DE LA CONGREGACIÓN]"
        markers_to_clear = tuple(
            marker for marker in TEMPLATE_MARKERS if marker != congregation_marker
        )
        cells_by_identity: dict[object, set[int]] = {}
        for table in self.document.tables:
            for row_index, row in enumerate(table.rows):
                for cell in row.cells:
                    cells_by_identity.setdefault(cell._tc, set()).add(row_index)

        found: set[str] = set()
        for table in self.document.tables:
            ranges = unused_ranges.get(table._tbl, [])
            seen_cells: set[object] = set()
            for row_index, row in enumerate(table.rows):
                time_cell = row.cells[0]._tc if row.cells else None
                for cell in row.cells:
                    if cell._tc in seen_cells:
                        continue
                    seen_cells.add(cell._tc)
                    cell_rows = cells_by_identity[cell._tc]
                    is_unused = bool(cell_rows) and all(
                        any(start <= cell_row < end for start, end in ranges)
                        for cell_row in cell_rows
                    )
                    text = cell.text
                    if congregation_marker in text:
                        self._replace_in_cell(
                            cell,
                            congregation_marker,
                            CONGREGATION_NAME,
                        )
                        text = cell.text
                    if is_unused:
                        for marker in markers_to_clear:
                            if marker in text:
                                self._remove_text(cell, marker)
                                text = cell.text
                        if cell._tc == time_cell and text.strip() == "0:00":
                            self._remove_text(cell, "0:00")
                            text = cell.text
                    found.update(marker for marker in TEMPLATE_MARKERS if marker in text)
                    if cell._tc == time_cell and text.strip() == "0:00":
                        found.add("0:00")
        return sorted(found)

    # ==================================================================
    # Reemplazo de texto preservando formato (nivel de runs)
    # ==================================================================

    def _set_cell_text(self, table, row_idx: int, col_idx: int, new_text: str,
                        old_text: str = "") -> None:
        if row_idx >= len(table.rows) or col_idx >= len(table.rows[row_idx].cells):
            return
        cell = table.rows[row_idx].cells[col_idx]
        search = old_text if old_text else cell.text.strip()
        if search and search in cell.text:
            self._replace_in_cell(cell, search, new_text)

    def _replace_in_cell_conditional(self, cell, old: str, new: str) -> bool:
        """Reemplaza solo si old está presente y new no está vacío."""
        if not new or old not in cell.text:
            return False
        return self._replace_in_cell(cell, old, new)

    def _paragraph_text(self, par) -> str:
        return "".join(run.text for run in par.runs)

    def _leading_number_prefix(self, text: str) -> str:
        stripped = text.lstrip()
        digits: list[str] = []
        for char in stripped:
            if char.isdigit():
                digits.append(char)
                continue
            break
        if not digits:
            return ""
        remaining = stripped[len(digits):].lstrip()
        if not remaining.startswith("."):
            return ""
        return f"{''.join(digits)}."

    def _first_duration_marker_index(self, text: str) -> int:
        markers = ("[X mins.]", "[XX mins.]", "(X mins.)", "(XX mins.)")
        positions = [text.find(marker) for marker in markers if text.find(marker) != -1]
        if not positions:
            return -1
        return min(positions)

    def _replace_in_cell(self, cell, old: str, new: str) -> bool:
        if not new:
            return False
        for par in cell.paragraphs:
            if self._replace_in_paragraph(par, old, new):
                return True
        return False

    def _remove_text(self, cell, text: str) -> None:
        for par in cell.paragraphs:
            self._replace_in_paragraph(par, text, "")

    def _replace_in_paragraph(self, par, search: str, repl: str) -> bool:
        runs = par.runs
        if not runs:
            return False
        full, b = "", []
        for ri, r in enumerate(runs):
            if r.text:
                s = len(full)
                full += r.text
                b.append((ri, s, len(full)))
        pos = full.find(search)
        if pos == -1:
            return False
        end = pos + len(search)
        aff = []
        for ri, rs, run_end in b:
            if rs <= pos < run_end:
                aff.append(ri)
            elif pos <= rs < end:
                aff.append(ri)
            elif rs < end <= run_end and ri not in aff:
                aff.append(ri)
        if not aff:
            return False
        aff.sort()
        f = aff[0]
        fs = next(s for r, s, e in b if r == f)
        rel = pos - fs
        orig = runs[f].text
        runs[f].text = orig[:rel] + repl + orig[rel + len(search):]
        for ri in aff[1:]:
            rs = next(s for r, s, e in b if r == ri)
            run_end = next(e for r, s, e in b if r == ri)
            if rs >= pos and run_end <= end:
                runs[ri].text = ""
            elif rs < end <= run_end:
                ov = end - rs
                if 0 < ov <= len(runs[ri].text):
                    runs[ri].text = runs[ri].text[ov:]
            elif rs <= pos < run_end:
                ov = run_end - pos
                if 0 < ov <= len(runs[ri].text):
                    runs[ri].text = runs[ri].text[:-ov]
        return True


def fill_template(
    template_path: Path,
    weeks: list[MeetingWeek],
    output_path: Path,
) -> GenerationResult:
    writer = TemplateWriter(template_path)
    return writer.fill(weeks, output_path)
